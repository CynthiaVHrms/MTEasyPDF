from __future__ import annotations

import re

from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from app.protocols_escudo_c5.document_editor import (
    extract_tag_key,
    read_tag_groups_from_excel,
)


TAG_PATTERN = re.compile(
    r"\{\{\s*([A-Za-z0-9ÁÉÍÓÚÜÑáéíóúüñ_-]+)\s*\}\}"
)

PENDING_PATTERN = re.compile(
    r"\bpendiente(?:s)?\b",
    flags=re.IGNORECASE,
)


@dataclass(frozen=True)
class TagOccurrence:
    tag: str
    paragraph_number: int
    paragraph_text: str


@dataclass(frozen=True)
class DuplicateTag:
    tag: str
    occurrences: int


@dataclass(frozen=True)
class TagValidationResult:
    compatible: bool
    word_tags: list[str]
    excel_tags: list[str]
    matched_tags: list[str]
    missing_in_word: list[str]
    missing_in_excel: list[str]
    duplicated_in_word: list[DuplicateTag]
    generic_pending_texts: list[str]
    total_word_tags: int
    total_excel_tags: int
    total_matches: int

    def to_dict(self) -> dict[str, object]:
        result = asdict(self)

        result["duplicated_in_word"] = [
            asdict(item)
            for item in self.duplicated_in_word
        ]

        return result


class InvalidTaggedWordError(ValueError):
    """Error controlado para documentos Word inválidos."""


def normalize_tag_key(value: object) -> str:
    """
    Convierte una etiqueta a la forma interna utilizada para comparar.

    Ejemplos:

    {{PENDIENTE_LIBERTAD}}
    pendiente_libertad
    PENDIENTE_LIBERTAD

    Resultado:

    PENDIENTE_LIBERTAD
    """

    tag_key = extract_tag_key(value)

    tag_key = re.sub(
        r"\s+",
        "_",
        tag_key.strip(),
    )

    tag_key = re.sub(
        r"[^A-Za-z0-9ÁÉÍÓÚÜÑáéíóúüñ_-]",
        "",
        tag_key,
    )

    return tag_key.upper()


def iter_document_paragraphs(
    container: DocumentObject | _Cell,
):
    """
    Recorre párrafos normales y párrafos que estén dentro de tablas.

    Esta búsqueda coincide con las zonas que actualmente puede modificar
    document_editor.py.
    """

    for paragraph in container.paragraphs:
        yield paragraph

    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_document_paragraphs(cell)


def extract_tags_from_paragraph(
    paragraph: Paragraph,
) -> list[str]:
    text = paragraph.text or ""

    return [
        normalize_tag_key(match.group(1))
        for match in TAG_PATTERN.finditer(text)
        if normalize_tag_key(match.group(1))
    ]


def read_word_tag_occurrences(
    document_path: Path,
) -> tuple[list[TagOccurrence], list[str]]:
    """
    Obtiene todas las etiquetas visibles dentro del cuerpo del Word.

    También detecta textos genéricos que contienen la palabra "Pendiente".
    Estos textos solamente se informan; nunca se reemplazan automáticamente.
    """

    document_path = Path(document_path)

    if not document_path.exists():
        raise FileNotFoundError(
            f"No existe el documento Word: {document_path}"
        )

    if not document_path.is_file():
        raise InvalidTaggedWordError(
            "La ruta del documento Word no corresponde a un archivo."
        )

    if document_path.suffix.lower() != ".docx":
        raise InvalidTaggedWordError(
            "El documento debe ser un archivo Word con extensión .docx."
        )

    try:
        document = Document(str(document_path))
    except Exception as exc:
        raise InvalidTaggedWordError(
            "No fue posible abrir el documento Word. "
            "Verifica que el archivo no esté dañado y que sea un .docx válido."
        ) from exc

    occurrences: list[TagOccurrence] = []
    generic_pending_texts: list[str] = []

    for paragraph_number, paragraph in enumerate(
        iter_document_paragraphs(document),
        start=1,
    ):
        paragraph_text = (paragraph.text or "").strip()

        if not paragraph_text:
            continue

        tags = extract_tags_from_paragraph(paragraph)

        for tag in tags:
            occurrences.append(
                TagOccurrence(
                    tag=tag,
                    paragraph_number=paragraph_number,
                    paragraph_text=paragraph_text,
                )
            )

        if PENDING_PATTERN.search(paragraph_text) and not tags:
            if paragraph_text not in generic_pending_texts:
                generic_pending_texts.append(paragraph_text)

    return occurrences, generic_pending_texts


def read_excel_tag_keys(
    excel_path: Path,
) -> list[str]:
    groups = read_tag_groups_from_excel(Path(excel_path))

    return sorted(
        {
            normalize_tag_key(group.tag_key)
            for group in groups
            if normalize_tag_key(group.tag_key)
        }
    )


def validate_word_and_excel_tags(
    document_path: Path,
    excel_path: Path,
) -> TagValidationResult:
    """
    Compara las etiquetas del Word contra las etiquetas del Excel.

    Reglas para considerar los archivos compatibles:

    1. Toda etiqueta del Excel debe aparecer en el Word.
    2. Toda etiqueta del Word debe existir en el Excel.
    3. Una etiqueta no debe aparecer más de una vez dentro del Word.
    """

    occurrences, generic_pending_texts = read_word_tag_occurrences(
        document_path=document_path,
    )

    excel_tags = read_excel_tag_keys(
        excel_path=excel_path,
    )

    occurrence_count: dict[str, int] = {}

    for occurrence in occurrences:
        occurrence_count[occurrence.tag] = (
            occurrence_count.get(occurrence.tag, 0) + 1
        )

    word_tags = sorted(occurrence_count.keys())
    excel_tag_set = set(excel_tags)
    word_tag_set = set(word_tags)

    matched_tags = sorted(
        word_tag_set.intersection(excel_tag_set)
    )

    missing_in_word = sorted(
        excel_tag_set.difference(word_tag_set)
    )

    missing_in_excel = sorted(
        word_tag_set.difference(excel_tag_set)
    )

    duplicated_in_word = sorted(
        [
            DuplicateTag(
                tag=tag,
                occurrences=count,
            )
            for tag, count in occurrence_count.items()
            if count > 1
        ],
        key=lambda item: item.tag,
    )

    compatible = (
        len(word_tags) > 0
        and len(excel_tags) > 0
        and not missing_in_word
        and not missing_in_excel
        and not duplicated_in_word
    )

    return TagValidationResult(
        compatible=compatible,
        word_tags=word_tags,
        excel_tags=excel_tags,
        matched_tags=matched_tags,
        missing_in_word=missing_in_word,
        missing_in_excel=missing_in_excel,
        duplicated_in_word=duplicated_in_word,
        generic_pending_texts=generic_pending_texts,
        total_word_tags=len(word_tags),
        total_excel_tags=len(excel_tags),
        total_matches=len(matched_tags),
    )


def format_validation_error(
    result: TagValidationResult,
) -> str:
    """
    Construye un mensaje entendible para el usuario final.
    """

    problems: list[str] = []

    if not result.word_tags:
        problems.append(
            "El Word no contiene etiquetas con el formato "
            "{{NOMBRE_DE_LA_ETIQUETA}}."
        )

    if result.missing_in_word:
        problems.append(
            "Etiquetas capturadas en el Excel que no aparecen en el Word: "
            + ", ".join(
                f"{{{{{tag}}}}}"
                for tag in result.missing_in_word
            )
            + "."
        )

    if result.missing_in_excel:
        problems.append(
            "Etiquetas encontradas en el Word que no existen en el Excel: "
            + ", ".join(
                f"{{{{{tag}}}}}"
                for tag in result.missing_in_excel
            )
            + "."
        )

    if result.duplicated_in_word:
        problems.append(
            "Existen etiquetas repetidas dentro del Word: "
            + ", ".join(
                (
                    f"{{{{{item.tag}}}}} "
                    f"({item.occurrences} veces)"
                )
                for item in result.duplicated_in_word
            )
            + ". Cada etiqueta debe identificar un único punto de inserción."
        )

    if not problems:
        return (
            "El Word y el Excel no pudieron validarse. "
            "Revisa las etiquetas capturadas."
        )

    return " ".join(problems)
from __future__ import annotations

import datetime
import json
import logging
import os
import re
import shutil
import stat
import unicodedata
import zipfile
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path, PurePosixPath
from typing import Callable
from uuid import uuid4

import pandas as pd
from PIL import Image, ImageFile, ImageOps, UnidentifiedImageError
from docx import Document
from docx.shared import Mm
from docxcompose.composer import Composer
from docxtpl import DocxTemplate, InlineImage


logger = logging.getLogger(__name__)

ImageFile.LOAD_TRUNCATED_IMAGES = True

IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".tif",
    ".tiff",
    ".webp",
    ".jfif",
}

WINDOWS_RESERVED_NAMES = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}


CLASSIFICATION_STOPWORDS = {
    "de",
    "del",
    "la",
    "las",
    "el",
    "los",
    "y",
    "en",
    "para",
}

CLASSIFICATION_TOKEN_REPLACEMENTS = {
    "equipos": "equipo",
    "herrajes": "herraje",
    "soportes": "soporte",
    "conectores": "conector",
    "numeros": "numero",
    "series": "serie",
    "etiquetas": "etiquetado",
    "etiquetada": "etiquetado",
    "etiquetado": "etiquetado",
    "rotulacion": "rotulado",
    "rotulada": "rotulado",
}


class InvalidEvidenceImageError(ValueError):
    """Error controlado para evidencias que no pueden convertirse a imagen."""


@dataclass
class ProtocolDocumentResult:
    job_id: str
    output_docx_path: Path
    total_sitios: int
    total_documentos_generados: int
    total_imagenes_omitidas: int = 0
    total_imagenes_no_encontradas: int = 0
    total_clasificaciones_no_encontradas: int = 0
    diagnostics_path: Path | None = None


ProtocolProgressCallback = Callable[[dict[str, object]], None]


def _emit_progress(
    callback: ProtocolProgressCallback | None,
    *,
    stage: str,
    percentage: int,
    message: str,
    current_site: int = 0,
    total_sites: int = 0,
    current_site_name: str = "",
    processed_images: int = 0,
    total_images: int = 0,
    detected_images: int = 0,
) -> None:
    """Publica avance sin permitir que un fallo informativo detenga el Word."""

    if callback is None:
        return

    payload: dict[str, object] = {
        "stage": stage,
        "percentage": max(0, min(100, int(percentage))),
        "message": message,
        "current_site": max(0, int(current_site)),
        "total_sites": max(0, int(total_sites)),
        "current_site_name": current_site_name,
        "processed_images": max(0, int(processed_images)),
        "total_images": max(0, int(total_images)),
        "detected_images": max(0, int(detected_images)),
    }

    try:
        callback(payload)
    except Exception:
        logger.exception(
            "No fue posible publicar el avance del trabajo. "
            "La generación continuará."
        )


def normalizar_texto(texto: object) -> str:
    """Normaliza texto para comparar nombres ignorando acentos y separadores."""

    if pd.isna(texto):
        return ""

    texto_normalizado = str(texto)
    texto_normalizado = "".join(
        caracter
        for caracter in unicodedata.normalize("NFD", texto_normalizado)
        if unicodedata.category(caracter) != "Mn"
    )
    texto_normalizado = texto_normalizado.lower()
    texto_normalizado = texto_normalizado.replace("-", " ").replace("_", " ")
    texto_normalizado = re.sub(r"[^a-z0-9 ]", "", texto_normalizado)
    texto_normalizado = re.sub(r"\s+", " ", texto_normalizado).strip()

    return texto_normalizado


def limpiar_prefijo_numerico(nombre_carpeta: object) -> str:
    """Elimina prefijos como ``1. `` o ``25. `` de un nombre de carpeta."""

    return re.sub(r"^\d+\.\s+", "", str(nombre_carpeta))


def _classification_tokens(value: object) -> set[str]:
    """Devuelve tokens comparables para nombres de clasificación."""

    normalized = normalizar_texto(limpiar_prefijo_numerico(value))
    tokens: set[str] = set()

    for token in normalized.split():
        token = CLASSIFICATION_TOKEN_REPLACEMENTS.get(token, token)

        if token and token not in CLASSIFICATION_STOPWORDS:
            tokens.add(token)

    return tokens


def canonicalizar_clasificacion(value: object) -> str:
    """
    Agrupa variantes frecuentes de una misma clasificación.

    Ejemplos equivalentes:
    - ``EQUIPO EN INDOOR`` / ``Equipo indoor``
    - ``NÚMERO DE SERIE`` / ``Números de serie``
    - ``ROTULADO`` / ``Etiquetado`` / ``Rotulado y etiquetado``
    - ``Condición de herrajes y soportes`` /
      ``Condición de soporte y herraje``
    """

    tokens = _classification_tokens(value)

    if "outdoor" in tokens:
        return "equipo outdoor"

    if "indoor" in tokens:
        return "equipo indoor"

    if "conector" in tokens:
        return "condicion conector"

    if "marca" in tokens and "modelo" in tokens and "serie" in tokens:
        return "marca modelo serie"

    if "serie" in tokens:
        return "numero serie"

    if "rotulado" in tokens or "etiquetado" in tokens:
        return "rotulado etiquetado"

    if "herraje" in tokens and "soporte" in tokens:
        return "condicion herraje soporte"

    if "herraje" in tokens:
        return "condicion herraje"

    if "soporte" in tokens:
        return "condicion soporte"

    return " ".join(sorted(tokens))


def _name_similarity(left: object, right: object) -> float:
    """Calcula una similitud conservadora entre dos nombres de carpeta."""

    left_normalized = normalizar_texto(limpiar_prefijo_numerico(left))
    right_normalized = normalizar_texto(limpiar_prefijo_numerico(right))

    if not left_normalized or not right_normalized:
        return 0.0

    if left_normalized == right_normalized:
        return 1.0

    left_tokens = _classification_tokens(left)
    right_tokens = _classification_tokens(right)

    if left_tokens or right_tokens:
        union = left_tokens | right_tokens
        token_score = len(left_tokens & right_tokens) / len(union)
    else:
        token_score = 0.0

    sequence_score = SequenceMatcher(
        None,
        left_normalized,
        right_normalized,
    ).ratio()

    return max(token_score, sequence_score)


def buscar_carpeta_clasificacion(
    ruta_sitio: str | Path | None,
    clasificacion_buscada: object,
) -> str | None:
    """
    Busca una clasificación con varias estrategias, sin depender de que el
    nombre sea idéntico al del Excel.

    El orden es:
    1. Coincidencia flexible exacta.
    2. Coincidencia canónica de clasificaciones conocidas.
    3. Coincidencia aproximada con umbral conservador.
    """

    exact_match = buscar_carpeta_flexible(
        ruta_padre=ruta_sitio,
        nombre_buscado=clasificacion_buscada,
    )

    if exact_match:
        return exact_match

    if not ruta_sitio or not _fs_is_dir(ruta_sitio):
        return None

    expected_canonical = canonicalizar_clasificacion(clasificacion_buscada)
    candidates = [
        child
        for child in _iter_directory(ruta_sitio)
        if _fs_is_dir(child)
    ]

    canonical_matches = [
        candidate
        for candidate in candidates
        if expected_canonical
        and canonicalizar_clasificacion(candidate.name) == expected_canonical
    ]

    if len(canonical_matches) == 1:
        logger.info(
            "Clasificación asociada por equivalencia: '%s' -> '%s'.",
            clasificacion_buscada,
            canonical_matches[0].name,
        )
        return str(canonical_matches[0])

    scored_candidates = sorted(
        (
            (_name_similarity(clasificacion_buscada, candidate.name), candidate)
            for candidate in candidates
        ),
        key=lambda item: item[0],
        reverse=True,
    )

    if not scored_candidates:
        return None

    best_score, best_candidate = scored_candidates[0]
    second_score = scored_candidates[1][0] if len(scored_candidates) > 1 else 0.0

    if best_score >= 0.62 and (best_score - second_score >= 0.08 or best_score >= 0.90):
        logger.info(
            "Clasificación asociada aproximadamente: '%s' -> '%s' (%.2f).",
            clasificacion_buscada,
            best_candidate.name,
            best_score,
        )
        return str(best_candidate)

    return None


def limpiar_fragmento_nombre_archivo(valor: object) -> str:
    """Genera un fragmento seguro para nombres de archivos temporales."""

    texto = str(valor).strip()
    texto = re.sub(r'[<>:"/\\|?*]', "_", texto)
    texto = re.sub(r"\s+", "_", texto)
    texto = texto.strip(" ._")

    return texto or "SIN_ID"


def texto_seguro_excel(valor: object) -> str:
    """Evita que valores vacíos de Excel se conviertan en la cadena ``nan``."""

    if pd.isna(valor):
        return ""

    texto = str(valor).strip()

    if texto.lower() == "nan":
        return ""

    return texto


def _to_extended_windows_path(path: str | Path) -> str:
    """
    Devuelve una ruta absoluta para operaciones de entrada/salida.

    En Windows agrega el prefijo ``\\?\\`` únicamente al momento de abrir,
    crear, recorrer o consultar archivos. El resto de la aplicación conserva
    rutas normales, porque ``pathlib`` y algunas librerías de terceros pueden
    perder la navegación de carpetas cuando reciben permanentemente una ruta
    con el prefijo extendido.
    """

    absolute_path = os.path.abspath(os.fspath(path))

    if os.name != "nt":
        return absolute_path

    if absolute_path.startswith("\\\\?\\"):
        return absolute_path

    if absolute_path.startswith("\\\\"):
        return "\\\\?\\UNC\\" + absolute_path[2:]

    return "\\\\?\\" + absolute_path

def _fs_exists(path: str | Path) -> bool:
    """Comprueba existencia usando rutas extendidas cuando se ejecuta en Windows."""

    return os.path.exists(_to_extended_windows_path(path))


def _fs_is_dir(path: str | Path) -> bool:
    """Comprueba si una ruta es carpeta sin perder compatibilidad con rutas largas."""

    return os.path.isdir(_to_extended_windows_path(path))


def _fs_is_file(path: str | Path) -> bool:
    """Comprueba si una ruta es archivo sin perder compatibilidad con rutas largas."""

    return os.path.isfile(_to_extended_windows_path(path))


def _iter_directory(path: str | Path) -> list[Path]:
    """
    Lista el contenido de una carpeta usando ``os.scandir`` sobre la ruta
    extendida, pero devuelve objetos ``Path`` normales para que la lógica de
    búsqueda y comparación continúe funcionando igual que antes.
    """

    normal_path = Path(path)

    if not _fs_is_dir(normal_path):
        return []

    children: list[Path] = []

    with os.scandir(_to_extended_windows_path(normal_path)) as entries:
        for entry in entries:
            children.append(normal_path / entry.name)

    return children


def _walk_directories(root: str | Path) -> list[Path]:
    """Recorre carpetas de forma compatible con rutas largas de Windows."""

    root_path = Path(root)
    directories: list[Path] = []
    pending: list[Path] = [root_path]

    while pending:
        current = pending.pop()

        for child in _iter_directory(current):
            if _fs_is_dir(child):
                directories.append(child)
                pending.append(child)

    return directories


def buscar_carpeta_flexible(
    ruta_padre: str | Path | None,
    nombre_buscado: object,
) -> str | None:
    """
    Busca una carpeta ignorando prefijos numéricos, acentos, mayúsculas,
    guiones y guiones bajos.

    La navegación se realiza con las funciones de compatibilidad de rutas
    largas para que el prefijo extendido de Windows no provoque falsos "no existe".
    """

    if not ruta_padre:
        return None

    ruta_padre_path = Path(ruta_padre)

    if not _fs_is_dir(ruta_padre_path):
        return None

    nombre_buscado_limpio = limpiar_prefijo_numerico(nombre_buscado)
    nombre_buscado_normalizado = normalizar_texto(nombre_buscado_limpio)

    if not nombre_buscado_normalizado:
        return None

    for carpeta in _iter_directory(ruta_padre_path):
        if not _fs_is_dir(carpeta):
            continue

        carpeta_limpia = limpiar_prefijo_numerico(carpeta.name)

        if normalizar_texto(carpeta_limpia) == nombre_buscado_normalizado:
            return str(carpeta)

    return None

def _zip_member_parts(member_name: str) -> tuple[str, ...]:
    """Valida y convierte una ruta interna del ZIP en componentes seguros."""

    normalized_name = member_name.replace("\\", "/")

    if "\x00" in normalized_name:
        raise ValueError("El ZIP contiene un nombre de archivo no permitido.")

    pure_path = PurePosixPath(normalized_name)

    if pure_path.is_absolute():
        raise ValueError("El ZIP contiene rutas absolutas no permitidas.")

    parts = tuple(part for part in pure_path.parts if part not in ("", "."))

    if not parts:
        return ()

    if any(part == ".." for part in parts):
        raise ValueError("El ZIP contiene rutas que intentan salir del directorio.")

    for part in parts:
        if len(part) > 255:
            raise ValueError(
                "El ZIP contiene un nombre individual mayor a 255 caracteres: "
                f"{part[:80]}..."
            )

        if any(character in part for character in '<>:"|?*'):
            raise ValueError(
                "El ZIP contiene caracteres no permitidos por Windows en: "
                f"{part}"
            )

        if any(ord(character) < 32 for character in part):
            raise ValueError(
                "El ZIP contiene caracteres de control no permitidos en: "
                f"{part}"
            )

        if part.endswith((" ", ".")):
            raise ValueError(
                "El ZIP contiene un nombre que termina en espacio o punto: "
                f"{part}"
            )

        base_name = part.split(".", 1)[0].upper()
        if base_name in WINDOWS_RESERVED_NAMES:
            raise ValueError(
                "El ZIP contiene un nombre reservado por Windows: "
                f"{part}"
            )

    return parts


def _is_zip_symlink(member: zipfile.ZipInfo) -> bool:
    """Detecta enlaces simbólicos almacenados en un ZIP creado en Unix."""

    unix_mode = (member.external_attr >> 16) & 0xFFFF
    return stat.S_ISLNK(unix_mode)


def safe_extract_zip(zip_path: Path, extract_dir: Path) -> Path:
    """
    Extrae un ZIP de forma segura y compatible con rutas largas de Windows.

    La ruta que se devuelve es una ruta normal. El prefijo extendido de
    Windows se aplica únicamente en las operaciones de entrada/salida. Esto
    evita que la lógica posterior de búsqueda de carpetas pierda las rutas de
    las evidencias y termine mostrando ``[No encontrada]`` en el documento.
    """

    zip_path = Path(zip_path)
    extract_dir = Path(extract_dir)

    if not _fs_exists(zip_path):
        raise FileNotFoundError(f"No existe el archivo ZIP: {zip_path}")

    if not zipfile.is_zipfile(_to_extended_windows_path(zip_path)):
        raise ValueError("El archivo de evidencias no es un ZIP válido.")

    os.makedirs(_to_extended_windows_path(extract_dir), exist_ok=True)

    try:
        with zipfile.ZipFile(_to_extended_windows_path(zip_path), "r") as zip_ref:
            for member in zip_ref.infolist():
                if member.flag_bits & 0x1:
                    raise ValueError(
                        "El ZIP contiene archivos protegidos con contraseña."
                    )

                if _is_zip_symlink(member):
                    raise ValueError(
                        "El ZIP contiene enlaces simbólicos, los cuales no están permitidos."
                    )

                parts = _zip_member_parts(member.filename)

                if not parts:
                    continue

                destination = extract_dir.joinpath(*parts)
                destination_io = _to_extended_windows_path(destination)

                if member.is_dir() or member.filename.endswith(("/", "\\")):
                    os.makedirs(destination_io, exist_ok=True)
                    continue

                os.makedirs(
                    _to_extended_windows_path(destination.parent),
                    exist_ok=True,
                )

                with zip_ref.open(member, "r") as source_file:
                    with open(destination_io, "wb") as destination_file:
                        shutil.copyfileobj(
                            source_file,
                            destination_file,
                            length=1024 * 1024,
                        )

    except ValueError:
        raise

    except OSError as exc:
        if getattr(exc, "winerror", None) == 206:
            raise ValueError(
                "No fue posible extraer el ZIP porque contiene una ruta que "
                "excede los límites admitidos por el sistema de archivos. "
                "La aplicación intentó utilizar rutas extendidas; revise que "
                "ningún nombre individual supere 255 caracteres."
            ) from exc

        raise ValueError(f"No fue posible extraer el ZIP: {exc}") from exc

    except zipfile.BadZipFile as exc:
        raise ValueError(
            "El archivo de evidencias está dañado o no es un ZIP válido."
        ) from exc

    return extract_dir

def find_evidence_root(
    extract_dir: Path,
    expected_link_names: list[str] | tuple[str, ...] | None = None,
) -> Path:
    """
    Detecta la carpeta que contiene directamente los enlaces del proyecto.

    Primero compara los nombres de enlace registrados en el Excel contra las
    carpetas extraídas. Esto es más confiable que buscar cualquier carpeta que
    contenga la palabra "evidencia", porque dentro del ZIP puede haber carpetas
    intermedias o clasificaciones con nombres similares.
    """

    extract_dir = Path(extract_dir)

    if not _fs_is_dir(extract_dir):
        return extract_dir

    expected_normalized = {
        normalizar_texto(limpiar_prefijo_numerico(name))
        for name in (expected_link_names or [])
        if normalizar_texto(name)
    }

    candidate_directories = [extract_dir, *_walk_directories(extract_dir)]

    if expected_normalized:
        best_candidate: Path | None = None
        best_score = 0

        for candidate in candidate_directories:
            child_names = {
                normalizar_texto(limpiar_prefijo_numerico(child.name))
                for child in _iter_directory(candidate)
                if _fs_is_dir(child)
            }

            score = len(expected_normalized.intersection(child_names))

            if score > best_score:
                best_candidate = candidate
                best_score = score

        if best_candidate is not None and best_score > 0:
            logger.info(
                "Raíz de evidencias detectada por coincidencia con Excel: %s",
                best_candidate,
            )
            return best_candidate

    evidence_names = {
        "evidencia de los enlaces",
        "evidencias de los enlaces",
        "evidencia enlaces",
        "evidencias enlaces",
    }

    evidence_candidates = [
        path
        for path in candidate_directories
        if normalizar_texto(path.name) in evidence_names
    ]

    if evidence_candidates:
        evidence_candidates.sort(key=lambda item: len(item.parts))
        logger.info(
            "Raíz de evidencias detectada por nombre de carpeta: %s",
            evidence_candidates[0],
        )
        return evidence_candidates[0]

    current = extract_dir

    while True:
        child_directories = [
            child
            for child in _iter_directory(current)
            if _fs_is_dir(child)
        ]
        child_files = [
            child
            for child in _iter_directory(current)
            if _fs_is_file(child)
        ]

        if len(child_directories) == 1 and not child_files:
            current = child_directories[0]
            continue

        break

    logger.info("Raíz de evidencias utilizada: %s", current)
    return current

def format_excel_date(value: object) -> str:
    """Convierte fechas de Excel al formato ``dd/mm/aaaa``."""

    if isinstance(value, (pd.Timestamp, datetime.datetime, datetime.date)):
        return value.strftime("%d/%m/%Y")

    if pd.isna(value):
        return ""

    return str(value)


def _strip_supported_image_extension(value: str) -> str:
    """
    Elimina únicamente extensiones de imagen conocidas.

    Es importante no utilizar ``Path.suffix`` directamente para valores como
    ``1.0`` porque Python interpreta ``.0`` como si fuera una extensión. Para
    este sistema, ``1`` y ``1.0`` representan el mismo identificador de foto.
    """

    text = value.strip()
    lower_text = text.casefold()

    for extension in sorted(IMAGE_EXTENSIONS, key=len, reverse=True):
        if lower_text.endswith(extension):
            return text[: -len(extension)]

    return text


def _normalizar_identificador_imagen(value: object) -> str:
    """
    Normaliza identificadores de imagen sin depender del sufijo ``.0``.

    Los siguientes valores se consideran equivalentes:

    - ``1``
    - ``01``
    - ``001``
    - ``1.0``
    - ``01.0``
    - ``1.jpg``
    - ``1.0.jpeg``
    - ``1 - evidencia.jpg``
    - ``1.0_foto.png``

    Para nombres no numéricos conserva una comparación textual normalizada.
    """

    text = texto_seguro_excel(value)

    if not text:
        return ""

    stem = _strip_supported_image_extension(Path(text).name).strip()

    numeric_match = re.fullmatch(r"0*(\d+)(?:[.,]0+)?", stem)

    if numeric_match:
        return str(int(numeric_match.group(1)))


    numeric_prefix_match = re.match(
        r"^\s*0*(\d+)(?:[.,]0+)?(?=$|[\s_-])",
        stem,
    )

    if numeric_prefix_match:
        return str(int(numeric_prefix_match.group(1)))

    return normalizar_texto(stem)

def _iter_files_recursive(
    root: str | Path | None,
    *,
    max_depth: int = 6,
) -> list[Path]:
    """Lista archivos de manera recursiva y compatible con rutas largas."""

    if not root or not _fs_is_dir(root):
        return []

    root_path = Path(root)
    files: list[Path] = []
    pending: list[tuple[Path, int]] = [(root_path, 0)]

    while pending:
        current, depth = pending.pop()

        for child in _iter_directory(current):
            if _fs_is_file(child):
                files.append(child)
            elif _fs_is_dir(child) and depth < max_depth:
                pending.append((child, depth + 1))

    return files


def _image_name_matches(file_path: Path, requested_name: object) -> bool:
    """
    Compara el nombre solicitado por el Excel con una imagen real.

    La comparación no depende de que el archivo se llame ``1.0.jpg``. También
    reconoce ``1.jpg``, ``01.jpeg``, ``001.png`` y nombres descriptivos que
    comienzan con el mismo identificador numérico.
    """

    if file_path.suffix.lower() not in IMAGE_EXTENSIONS:
        return False

    requested_text = texto_seguro_excel(requested_name)

    if not requested_text:
        return False

    requested_base = _strip_supported_image_extension(
        Path(requested_text).name
    ).strip()
    file_base = _strip_supported_image_extension(file_path.name).strip()

    if requested_base.casefold() == file_base.casefold():
        return True

    requested_id = _normalizar_identificador_imagen(requested_text)
    file_id = _normalizar_identificador_imagen(file_path.name)

    if requested_id and requested_id == file_id:
        return True

    
    if requested_id.isdigit():
        numeric_prefix = re.match(
            r"^\s*0*(\d+)(?:[.,]0+)?(?=$|[\s_-])",
            file_base,
            flags=re.IGNORECASE,
        )

        if numeric_prefix:
            return str(int(numeric_prefix.group(1))) == requested_id

    return False



def buscar_imagen_flexible(
    ruta_carpeta: str | Path | None,
    nombre_evidencia: object,
) -> Path | None:
    """
    Busca una imagen dentro de la clasificación y sus subcarpetas.

    Esta búsqueda acepta diferencias de extensión, mayúsculas, valores
    numéricos provenientes de Excel y carpetas intermedias como ``Fotos``.
    """

    if not ruta_carpeta or not _fs_is_dir(ruta_carpeta):
        return None

    matches = [
        file_path
        for file_path in _iter_files_recursive(ruta_carpeta)
        if _image_name_matches(file_path, nombre_evidencia)
    ]

    if len(matches) == 1:
        return matches[0]

    if len(matches) > 1:
        # Preferir el archivo menos profundo y con nombre más corto.
        matches.sort(
            key=lambda item: (
                len(item.relative_to(Path(ruta_carpeta)).parts),
                len(item.name),
                item.name.casefold(),
            )
        )
        return matches[0]

    return None


def buscar_imagen_contextual(
    *,
    ruta_clasificacion: str | Path | None,
    ruta_sitio: str | Path | None,
    ruta_enlace: str | Path | None,
    clasificacion: object,
    nombre_evidencia: object,
) -> Path | None:
    """
    Localiza una evidencia aunque el ZIP tenga una carpeta extra o una
    variante razonable en el nombre de la clasificación.

    Nunca selecciona una coincidencia global ambigua. La prioridad es:
    1. Carpeta de clasificación resuelta.
    2. Candidatos bajo el sitio, puntuados por su carpeta padre.
    3. Candidato único bajo el enlace.
    """

    direct_match = buscar_imagen_flexible(
        ruta_carpeta=ruta_clasificacion,
        nombre_evidencia=nombre_evidencia,
    )

    if direct_match:
        return direct_match

    site_candidates = [
        file_path
        for file_path in _iter_files_recursive(ruta_sitio)
        if _image_name_matches(file_path, nombre_evidencia)
    ]

    if len(site_candidates) == 1:
        logger.info(
            "Imagen '%s' encontrada mediante búsqueda recursiva del sitio: %s",
            nombre_evidencia,
            site_candidates[0],
        )
        return site_candidates[0]

    if len(site_candidates) > 1:
        expected_canonical = canonicalizar_clasificacion(clasificacion)
        scored: list[tuple[float, Path]] = []

        for candidate in site_candidates:
            # Se revisan hasta tres carpetas ascendentes porque algunas
            # estructuras agregan carpetas intermedias como ``Fotos``.
            ancestors = list(candidate.parents)[:3]
            scores = []

            for ancestor in ancestors:
                canonical_score = (
                    1.0
                    if expected_canonical
                    and canonicalizar_clasificacion(ancestor.name) == expected_canonical
                    else 0.0
                )
                scores.append(
                    max(
                        canonical_score,
                        _name_similarity(clasificacion, ancestor.name),
                    )
                )

            scored.append((max(scores, default=0.0), candidate))

        scored.sort(key=lambda item: item[0], reverse=True)
        best_score, best_candidate = scored[0]
        second_score = scored[1][0] if len(scored) > 1 else 0.0

        if best_score >= 0.62 and (best_score - second_score >= 0.08 or best_score >= 0.95):
            logger.info(
                "Imagen '%s' asociada por contexto de clasificación '%s': %s",
                nombre_evidencia,
                clasificacion,
                best_candidate,
            )
            return best_candidate

    link_candidates = [
        file_path
        for file_path in _iter_files_recursive(ruta_enlace)
        if _image_name_matches(file_path, nombre_evidencia)
    ]

    if len(link_candidates) == 1:
        logger.info(
            "Imagen '%s' encontrada como coincidencia única del enlace: %s",
            nombre_evidencia,
            link_candidates[0],
        )
        return link_candidates[0]

    return None

def normalize_image_for_docx(
    source_path: Path,
    normalized_images_dir: Path,
) -> Path:
    """
    Corrige orientación, elimina metadatos problemáticos y genera un PNG limpio.
    """

    source_path = Path(source_path)
    normalized_images_dir = Path(normalized_images_dir)

    if not _fs_exists(source_path):
        raise InvalidEvidenceImageError(
            f"No se encontró la imagen: {source_path.name}"
        )

    if not _fs_is_file(source_path):
        raise InvalidEvidenceImageError(
            f"La ruta no corresponde a una imagen: {source_path}"
        )

    normalized_images_dir.mkdir(parents=True, exist_ok=True)
    normalized_path = normalized_images_dir / (
        f"{limpiar_fragmento_nombre_archivo(source_path.stem)}_"
        f"{uuid4().hex}.png"
    )

    try:
        with Image.open(_to_extended_windows_path(source_path)) as source_image:
            source_image.load()
            corrected_image = ImageOps.exif_transpose(source_image)

            if corrected_image.mode in ("RGBA", "LA"):
                normalized_image = corrected_image.convert("RGBA")
            elif corrected_image.mode == "P":
                if "transparency" in corrected_image.info:
                    normalized_image = corrected_image.convert("RGBA")
                else:
                    normalized_image = corrected_image.convert("RGB")
            else:
                normalized_image = corrected_image.convert("RGB")

            normalized_image.save(
                normalized_path,
                format="PNG",
                optimize=True,
            )

    except UnidentifiedImageError as exc:
        raise InvalidEvidenceImageError(
            "El archivo no contiene una imagen válida: "
            f"{source_path.name}"
        ) from exc

    except (OSError, ValueError, SyntaxError) as exc:
        raise InvalidEvidenceImageError(
            "La imagen está dañada, incompleta o utiliza un formato "
            f"incompatible: {source_path.name}"
        ) from exc

    if not normalized_path.exists() or normalized_path.stat().st_size == 0:
        raise InvalidEvidenceImageError(
            f"No fue posible preparar la imagen: {source_path.name}"
        )

    return normalized_path


def build_inline_image(
    template: DocxTemplate,
    image_path: str | Path,
    normalized_images_dir: Path,
) -> InlineImage:
    """Normaliza una imagen y crea el ``InlineImage`` para docxtpl."""

    normalized_path = normalize_image_for_docx(
        source_path=Path(image_path),
        normalized_images_dir=normalized_images_dir,
    )

    try:
        with Image.open(normalized_path) as image:
            width_pixels, height_pixels = image.size
    except Exception as exc:
        raise InvalidEvidenceImageError(
            "No fue posible leer la imagen normalizada: "
            f"{Path(image_path).name}"
        ) from exc

    if width_pixels >= height_pixels:
        return InlineImage(template, str(normalized_path), width=Mm(70))

    return InlineImage(template, str(normalized_path), height=Mm(52.5))


def generate_protocol_document(
    excel_path: Path,
    evidence_zip_path: Path,
    template_docx_path: Path,
    work_dir: Path,
    *,
    job_id: str | None = None,
    progress_callback: ProtocolProgressCallback | None = None,
) -> ProtocolDocumentResult:
    """Genera el protocolo Word final a partir del Excel y ZIP de evidencias."""

    _emit_progress(
        progress_callback,
        stage="validating_files",
        percentage=3,
        message="Validando los archivos recibidos.",
    )

    excel_path = Path(excel_path)
    evidence_zip_path = Path(evidence_zip_path)
    template_docx_path = Path(template_docx_path)
    work_dir = Path(work_dir)

    if not excel_path.exists():
        raise FileNotFoundError(f"No existe el Excel: {excel_path}")

    if not evidence_zip_path.exists():
        raise FileNotFoundError(
            f"No existe el ZIP de evidencias: {evidence_zip_path}"
        )

    if not template_docx_path.exists():
        raise FileNotFoundError(
            f"No existe la plantilla Word: {template_docx_path}"
        )

    job_id = job_id or uuid4().hex
    job_dir = work_dir / job_id

    # Nombres deliberadamente cortos. Además del soporte de rutas extendidas,
    # esto reduce la longitud total y mejora compatibilidad con herramientas
    # de terceros que todavía usan MAX_PATH.
    extract_dir = job_dir / "e"
    normalized_images_dir = job_dir / "i"
    temp_dir = job_dir / "t"
    output_dir = job_dir / "o"
    output_docx_path = output_dir / "Protocolo_C5_Generado.docx"

    if job_dir.exists():
        shutil.rmtree(job_dir, ignore_errors=True)

    extract_dir.mkdir(parents=True, exist_ok=True)
    normalized_images_dir.mkdir(parents=True, exist_ok=True)
    temp_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)

    _emit_progress(
        progress_callback,
        stage="preparing_workspace",
        percentage=6,
        message="Preparando el espacio de trabajo.",
    )
    _emit_progress(
        progress_callback,
        stage="extracting_zip",
        percentage=8,
        message="Extrayendo y validando el ZIP de evidencias.",
    )

    effective_extract_dir = safe_extract_zip(
        zip_path=evidence_zip_path,
        extract_dir=extract_dir,
    )

    _emit_progress(
        progress_callback,
        stage="extracting_zip",
        percentage=16,
        message="ZIP extraído correctamente.",
    )
    _emit_progress(
        progress_callback,
        stage="reading_excel",
        percentage=18,
        message="Leyendo y validando la plantilla Excel.",
    )

    try:
        hoja1 = pd.read_excel(excel_path, sheet_name=0)
        hoja2 = pd.read_excel(excel_path, sheet_name=1)
    except Exception as exc:
        raise ValueError(
            "No fue posible leer la plantilla Excel. Verifique que el archivo "
            "corresponda a la plantilla oficial y que no se encuentre dañado."
        ) from exc

    hoja1.columns = hoja1.columns.astype(str).str.strip()
    hoja2.columns = hoja2.columns.astype(str).str.strip()

    required_hoja1 = ["ID", "Sitio"]
    required_hoja2 = ["ID", "Sitio", "clasificacion", "Foto"]

    missing_hoja1 = [
        column for column in required_hoja1 if column not in hoja1.columns
    ]
    missing_hoja2 = [
        column for column in required_hoja2 if column not in hoja2.columns
    ]

    if missing_hoja1:
        raise ValueError(
            "La primera hoja del Excel no contiene las columnas requeridas: "
            + ", ".join(missing_hoja1)
        )

    if missing_hoja2:
        raise ValueError(
            "La segunda hoja del Excel no contiene las columnas requeridas: "
            + ", ".join(missing_hoja2)
        )

    hoja1["ID"] = hoja1["ID"].astype(str).str.strip()
    hoja2["ID"] = hoja2["ID"].astype(str).str.strip()

    _emit_progress(
        progress_callback,
        stage="reading_excel",
        percentage=22,
        message="La plantilla Excel fue validada correctamente.",
        total_sites=len(hoja1),
    )

    expected_link_names = [
        texto_seguro_excel(value)
        for value in hoja1.get("Enlace", pd.Series(dtype=object)).tolist()
        if texto_seguro_excel(value)
    ]

    ruta_evidencias = find_evidence_root(
        extract_dir=effective_extract_dir,
        expected_link_names=expected_link_names,
    )

    archivos_generados: list[Path] = []
    total_sitios = len(hoja1)
    total_imagenes_omitidas = 0
    total_imagenes_insertadas = 0
    total_imagenes_no_encontradas = 0
    total_clasificaciones_no_encontradas = 0

    imagenes_no_encontradas: list[dict[str, str]] = []
    imagenes_omitidas: list[dict[str, str]] = []
    clasificaciones_no_encontradas: list[dict[str, str]] = []
    clasificaciones_no_encontradas_keys: set[tuple[str, str]] = set()

    total_imagenes_en_zip = sum(
        1
        for file_path in _iter_files_recursive(ruta_evidencias, max_depth=12)
        if file_path.suffix.lower() in IMAGE_EXTENSIONS
    )
    total_imagenes_esperadas = sum(
        1
        for value in hoja2["Foto"].tolist()
        if texto_seguro_excel(value)
    )
    imagenes_procesadas = 0
    progress_update_step = max(1, total_imagenes_esperadas // 100)
    last_progress_percentage = -1

    def report_processing_progress(
        *,
        current_site: int,
        current_site_name: str,
        force: bool = False,
    ) -> None:
        nonlocal last_progress_percentage

        image_ratio = (
            imagenes_procesadas / total_imagenes_esperadas
            if total_imagenes_esperadas
            else 0.0
        )
        site_ratio = (
            max(0, current_site - 1) / total_sitios
            if total_sitios
            else 0.0
        )
        work_ratio = image_ratio if total_imagenes_esperadas else site_ratio
        percentage = 30 + int(min(1.0, work_ratio) * 56)

        if not force and percentage == last_progress_percentage:
            return

        last_progress_percentage = percentage
        _emit_progress(
            progress_callback,
            stage="processing_sites",
            percentage=percentage,
            message=(
                f"Procesando sitio {current_site} de {total_sitios}: "
                f"{current_site_name}"
            ),
            current_site=current_site,
            total_sites=total_sitios,
            current_site_name=current_site_name,
            processed_images=imagenes_procesadas,
            total_images=total_imagenes_esperadas,
            detected_images=total_imagenes_en_zip,
        )

    logger.info(
        "Diagnóstico del ZIP: raíz='%s', imágenes detectadas=%s.",
        ruta_evidencias,
        total_imagenes_en_zip,
    )

    if total_imagenes_en_zip == 0:
        raise ValueError(
            "El ZIP de evidencias no contiene imágenes compatibles. "
            "Primero extraiga el ZIP de estructura, coloque las evidencias "
            "dentro de las carpetas correspondientes y vuelva a comprimirlo."
        )

    if total_sitios == 0:
        raise ValueError("La primera hoja del Excel no contiene registros.")

    _emit_progress(
        progress_callback,
        stage="analyzing_evidence",
        percentage=28,
        message=(
            f"Se detectaron {total_imagenes_en_zip} imágenes en el ZIP. "
            "Iniciando la integración de evidencias."
        ),
        total_sites=total_sitios,
        total_images=total_imagenes_esperadas,
        detected_images=total_imagenes_en_zip,
    )
    _emit_progress(
        progress_callback,
        stage="processing_sites",
        percentage=30,
        message="Iniciando el procesamiento de sitios y evidencias.",
        total_sites=total_sitios,
        total_images=total_imagenes_esperadas,
        detected_images=total_imagenes_en_zip,
    )

    for index, fila_sitio in hoja1.iterrows():
        id_sitio = texto_seguro_excel(fila_sitio["ID"])
        nombre_sitio_real = texto_seguro_excel(fila_sitio["Sitio"])
        nombre_enlace = texto_seguro_excel(fila_sitio.get("Enlace", ""))
        site_number = index + 1
        report_processing_progress(
            current_site=site_number,
            current_site_name=nombre_sitio_real,
            force=True,
        )
        tipo_sitio_original = texto_seguro_excel(
            fila_sitio.get("TipoDeSitio", "")
        ).upper()

        if "LOCAL" in tipo_sitio_original:
            tipo_sitio_formateado = "Sitio Local"
        else:
            tipo_sitio_formateado = "Sitio Remoto"

        template = DocxTemplate(str(template_docx_path))

        registros_fotos = hoja2[
            (hoja2["ID"] == id_sitio)
            & (
                hoja2["Sitio"].astype(str).str.strip()
                == nombre_sitio_real
            )
        ]

        clasificaciones_unicas = (
            registros_fotos["clasificacion"]
            .fillna("SIN CLASIFICACION")
            .astype(str)
            .str.strip()
            .unique()
        )

        lista_clasificaciones = []

        ruta_enlace_folder = buscar_carpeta_flexible(
            ruta_padre=ruta_evidencias,
            nombre_buscado=nombre_enlace,
        )
        ruta_sitio_folder = buscar_carpeta_flexible(
            ruta_padre=ruta_enlace_folder,
            nombre_buscado=nombre_sitio_real,
        )

        if not ruta_enlace_folder:
            logger.warning(
                "No se encontró la carpeta del enlace '%s' dentro de '%s'.",
                nombre_enlace,
                ruta_evidencias,
            )

        elif not ruta_sitio_folder:
            logger.warning(
                "No se encontró la carpeta del sitio '%s' dentro del enlace '%s'.",
                nombre_sitio_real,
                ruta_enlace_folder,
            )

        for clasificacion in clasificaciones_unicas:
            fotos_clase = registros_fotos[
                (
                    registros_fotos["clasificacion"]
                    .fillna("SIN CLASIFICACION")
                    .astype(str)
                    .str.strip()
                )
                == clasificacion
            ]

            ruta_clasificacion = buscar_carpeta_clasificacion(
                ruta_sitio=ruta_sitio_folder,
                clasificacion_buscada=clasificacion,
            )

            if not ruta_clasificacion:
                logger.warning(
                    "No se encontró la clasificación '%s' para el sitio '%s'.",
                    clasificacion,
                    nombre_sitio_real,
                )

                classification_key = (
                    nombre_sitio_real.casefold(),
                    str(clasificacion).casefold(),
                )

                if classification_key not in clasificaciones_no_encontradas_keys:
                    clasificaciones_no_encontradas_keys.add(classification_key)
                    total_clasificaciones_no_encontradas += 1
                    clasificaciones_no_encontradas.append(
                        {
                            "sitio": nombre_sitio_real,
                            "clasificacion": str(clasificacion),
                            "enlace": nombre_enlace,
                        }
                    )

            lista_individual = []

            for _, fila_foto in fotos_clase.iterrows():
                nombre_evidencia = texto_seguro_excel(
                    fila_foto.get("Foto", "")
                )
                nombre_documento = texto_seguro_excel(
                    fila_foto.get("NombreDeImagenEnDocumento", "")
                )

                if not nombre_documento:
                    nombre_documento = " "

                datos_bloque = {
                    "nombre_foto": nombre_documento,
                    "tipo_equipo": texto_seguro_excel(
                        fila_foto.get("TipoDeEquipo", "")
                    ),
                    "sitio": nombre_sitio_real,
                    "imagen": "",
                }

                ruta_imagen = buscar_imagen_contextual(
                    ruta_clasificacion=ruta_clasificacion,
                    ruta_sitio=ruta_sitio_folder,
                    ruta_enlace=ruta_enlace_folder,
                    clasificacion=clasificacion,
                    nombre_evidencia=nombre_evidencia,
                )

                if ruta_imagen:
                    try:
                        datos_bloque["imagen"] = build_inline_image(
                            template=template,
                            image_path=ruta_imagen,
                            normalized_images_dir=normalized_images_dir,
                        )
                        total_imagenes_insertadas += 1
                    except InvalidEvidenceImageError as exc:
                        total_imagenes_omitidas += 1
                        logger.warning(
                            "La imagen '%s' fue omitida. Motivo: %s",
                            ruta_imagen,
                            exc,
                        )
                        imagenes_omitidas.append(
                            {
                                "sitio": nombre_sitio_real,
                                "clasificacion": str(clasificacion),
                                "imagen_solicitada": nombre_evidencia,
                                "archivo_detectado": ruta_imagen.name,
                                "motivo": str(exc),
                            }
                        )
                        datos_bloque["imagen"] = (
                            "[Imagen no disponible: "
                            f"{nombre_evidencia}]"
                        )
                else:
                    total_imagenes_no_encontradas += 1
                    logger.warning(
                        "No se encontró la imagen '%s' para sitio='%s', "
                        "clasificación='%s'. Carpeta de clasificación=%s.",
                        nombre_evidencia,
                        nombre_sitio_real,
                        clasificacion,
                        ruta_clasificacion,
                    )

                    imagenes_no_encontradas.append(
                        {
                            "sitio": nombre_sitio_real,
                            "clasificacion": str(clasificacion),
                            "imagen_solicitada": nombre_evidencia,
                            "enlace": nombre_enlace,
                            "motivo": (
                                "No se encontró la carpeta de clasificación."
                                if not ruta_clasificacion
                                else "No existe un archivo compatible dentro de la clasificación."
                            ),
                        }
                    )

                    datos_bloque["imagen"] = (
                        f"[No encontrada: {nombre_evidencia}]"
                    )

                lista_individual.append(datos_bloque)
                imagenes_procesadas += 1

                if (
                    imagenes_procesadas % progress_update_step == 0
                    or imagenes_procesadas == total_imagenes_esperadas
                ):
                    report_processing_progress(
                        current_site=site_number,
                        current_site_name=nombre_sitio_real,
                    )

            lista_pares = []

            for position in range(0, len(lista_individual), 2):
                pair = {
                    "izq": lista_individual[position],
                    "der": (
                        lista_individual[position + 1]
                        if position + 1 < len(lista_individual)
                        else None
                    ),
                }
                lista_pares.append(pair)

            lista_clasificaciones.append(
                {
                    "clasificacion": clasificacion,
                    "pares": lista_pares,
                }
            )

        fecha_arribo = format_excel_date(
            fila_sitio.get("FechaDeArribo", "")
        )
        fecha_terminacion = format_excel_date(
            fila_sitio.get("FechaDeTerminación", "")
        )

        datos_render = {
            "ID": id_sitio,
            "Enlace": nombre_enlace,
            "Sitio": nombre_sitio_real,
            "NoTicket": texto_seguro_excel(fila_sitio.get("NoTicket", "")),
            "TipoDeSitio": tipo_sitio_formateado,
            "TipoDeEquipo": texto_seguro_excel(
                fila_sitio.get("TipoDeEquipo", "")
            ),
            "Dato1": fecha_arribo,
            "Dato2": texto_seguro_excel(
                fila_sitio.get("HoraDeArribo", "")
            ),
            "Dato3": fecha_terminacion,
            "Dato4": texto_seguro_excel(
                fila_sitio.get("HoraDeTerminación", "")
            ),
            "clasificaciones": lista_clasificaciones,
        }

        template.render(datos_render)

        safe_id = limpiar_fragmento_nombre_archivo(id_sitio)
        temporal_path = temp_dir / f"temp_sitio_{safe_id}_{index}.docx"
        template.save(str(temporal_path))
        archivos_generados.append(temporal_path)
        report_processing_progress(
            current_site=site_number,
            current_site_name=nombre_sitio_real,
            force=True,
        )

    _emit_progress(
        progress_callback,
        stage="writing_diagnostics",
        percentage=88,
        message="Preparando el resumen de evidencias y observaciones.",
        current_site=total_sitios,
        total_sites=total_sitios,
        processed_images=imagenes_procesadas,
        total_images=total_imagenes_esperadas,
        detected_images=total_imagenes_en_zip,
    )

    logger.info(
        "Resumen de imágenes: detectadas_en_zip=%s, insertadas=%s, "
        "no_encontradas=%s, omitidas_por_daño=%s.",
        total_imagenes_en_zip,
        total_imagenes_insertadas,
        total_imagenes_no_encontradas,
        total_imagenes_omitidas,
    )

    diagnostics_path = output_dir / "diagnostico_generacion.json"
    diagnostics_payload = {
        "schema_version": "1.0",
        "job_id": job_id,
        "estado": "completado_con_observaciones" if (
            total_imagenes_no_encontradas
            or total_imagenes_omitidas
            or total_clasificaciones_no_encontradas
        ) else "completado",
        "resumen": {
            "imagenes_detectadas_en_zip": total_imagenes_en_zip,
            "imagenes_insertadas": total_imagenes_insertadas,
            "imagenes_no_encontradas": total_imagenes_no_encontradas,
            "imagenes_omitidas_por_dano": total_imagenes_omitidas,
            "clasificaciones_no_encontradas": total_clasificaciones_no_encontradas,
        },
        "imagenes_no_encontradas": imagenes_no_encontradas,
        "imagenes_omitidas": imagenes_omitidas,
        "clasificaciones_no_encontradas": clasificaciones_no_encontradas,
    }

    diagnostics_path.write_text(
        json.dumps(
            diagnostics_payload,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    _emit_progress(
        progress_callback,
        stage="writing_diagnostics",
        percentage=91,
        message="Resumen de evidencias generado.",
        current_site=total_sitios,
        total_sites=total_sitios,
        processed_images=imagenes_procesadas,
        total_images=total_imagenes_esperadas,
        detected_images=total_imagenes_en_zip,
    )

    if total_imagenes_insertadas == 0:
        raise ValueError(
            "El ZIP sí contiene imágenes, pero ninguna pudo asociarse con la "
            "plantilla Excel. Verifique que esté utilizando el mismo Excel "
            "con el que se generó la estructura y que las evidencias conserven "
            "los nombres indicados en la columna 'Foto'."
        )

    if not archivos_generados:
        raise ValueError("No se generaron documentos.")

    _emit_progress(
        progress_callback,
        stage="assembling_document",
        percentage=92,
        message="Uniendo los documentos de cada sitio.",
        current_site=total_sitios,
        total_sites=total_sitios,
        processed_images=imagenes_procesadas,
        total_images=total_imagenes_esperadas,
        detected_images=total_imagenes_en_zip,
    )

    documento_base = Document(str(archivos_generados[0]))
    composer = Composer(documento_base)
    total_archivos = len(archivos_generados)

    for document_index, archivo in enumerate(archivos_generados[1:], start=2):
        documento_base.add_page_break()
        documento_temporal = Document(str(archivo))
        composer.append(documento_temporal)

        assembly_percentage = 92 + int(
            (document_index / max(1, total_archivos)) * 5
        )
        _emit_progress(
            progress_callback,
            stage="assembling_document",
            percentage=min(97, assembly_percentage),
            message=(
                f"Uniendo documento {document_index} de {total_archivos}."
            ),
            current_site=total_sitios,
            total_sites=total_sitios,
            processed_images=imagenes_procesadas,
            total_images=total_imagenes_esperadas,
            detected_images=total_imagenes_en_zip,
        )

    _emit_progress(
        progress_callback,
        stage="saving_document",
        percentage=98,
        message="Guardando el documento Word final.",
        current_site=total_sitios,
        total_sites=total_sitios,
        processed_images=imagenes_procesadas,
        total_images=total_imagenes_esperadas,
        detected_images=total_imagenes_en_zip,
    )
    composer.save(str(output_docx_path))

    _emit_progress(
        progress_callback,
        stage="saving_document",
        percentage=99,
        message="Documento Word creado. Preparando la descarga.",
        current_site=total_sitios,
        total_sites=total_sitios,
        processed_images=imagenes_procesadas,
        total_images=total_imagenes_esperadas,
        detected_images=total_imagenes_en_zip,
    )

    if not output_docx_path.exists() or output_docx_path.stat().st_size == 0:
        raise RuntimeError("El documento Word no se generó correctamente.")

    return ProtocolDocumentResult(
        job_id=job_id,
        output_docx_path=output_docx_path,
        total_sitios=total_sitios,
        total_documentos_generados=len(archivos_generados),
        total_imagenes_omitidas=total_imagenes_omitidas,
        total_imagenes_no_encontradas=total_imagenes_no_encontradas,
        total_clasificaciones_no_encontradas=(
            total_clasificaciones_no_encontradas
        ),
        diagnostics_path=diagnostics_path,
    )
















# from __future__ import annotations

# import datetime
# import json
# import logging
# import os
# import re
# import shutil
# import stat
# import unicodedata
# import zipfile
# from dataclasses import dataclass
# from difflib import SequenceMatcher
# from pathlib import Path, PurePosixPath
# from uuid import uuid4

# import pandas as pd
# from PIL import Image, ImageFile, ImageOps, UnidentifiedImageError
# from docx import Document
# from docx.shared import Mm
# from docxcompose.composer import Composer
# from docxtpl import DocxTemplate, InlineImage


# logger = logging.getLogger(__name__)

# ImageFile.LOAD_TRUNCATED_IMAGES = True

# IMAGE_EXTENSIONS = {
#     ".jpg",
#     ".jpeg",
#     ".png",
#     ".bmp",
#     ".tif",
#     ".tiff",
#     ".webp",
#     ".jfif",
# }

# WINDOWS_RESERVED_NAMES = {
#     "CON",
#     "PRN",
#     "AUX",
#     "NUL",
#     *(f"COM{i}" for i in range(1, 10)),
#     *(f"LPT{i}" for i in range(1, 10)),
# }


# CLASSIFICATION_STOPWORDS = {
#     "de",
#     "del",
#     "la",
#     "las",
#     "el",
#     "los",
#     "y",
#     "en",
#     "para",
# }

# CLASSIFICATION_TOKEN_REPLACEMENTS = {
#     "equipos": "equipo",
#     "herrajes": "herraje",
#     "soportes": "soporte",
#     "conectores": "conector",
#     "numeros": "numero",
#     "series": "serie",
#     "etiquetas": "etiquetado",
#     "etiquetada": "etiquetado",
#     "etiquetado": "etiquetado",
#     "rotulacion": "rotulado",
#     "rotulada": "rotulado",
# }


# class InvalidEvidenceImageError(ValueError):
#     """Error controlado para evidencias que no pueden convertirse a imagen."""


# @dataclass
# class ProtocolDocumentResult:
#     job_id: str
#     output_docx_path: Path
#     total_sitios: int
#     total_documentos_generados: int
#     total_imagenes_omitidas: int = 0
#     total_imagenes_no_encontradas: int = 0
#     total_clasificaciones_no_encontradas: int = 0
#     diagnostics_path: Path | None = None


# def normalizar_texto(texto: object) -> str:
#     """Normaliza texto para comparar nombres ignorando acentos y separadores."""

#     if pd.isna(texto):
#         return ""

#     texto_normalizado = str(texto)
#     texto_normalizado = "".join(
#         caracter
#         for caracter in unicodedata.normalize("NFD", texto_normalizado)
#         if unicodedata.category(caracter) != "Mn"
#     )
#     texto_normalizado = texto_normalizado.lower()
#     texto_normalizado = texto_normalizado.replace("-", " ").replace("_", " ")
#     texto_normalizado = re.sub(r"[^a-z0-9 ]", "", texto_normalizado)
#     texto_normalizado = re.sub(r"\s+", " ", texto_normalizado).strip()

#     return texto_normalizado


# def limpiar_prefijo_numerico(nombre_carpeta: object) -> str:
#     """Elimina prefijos como ``1. `` o ``25. `` de un nombre de carpeta."""

#     return re.sub(r"^\d+\.\s+", "", str(nombre_carpeta))


# def _classification_tokens(value: object) -> set[str]:
#     """Devuelve tokens comparables para nombres de clasificación."""

#     normalized = normalizar_texto(limpiar_prefijo_numerico(value))
#     tokens: set[str] = set()

#     for token in normalized.split():
#         token = CLASSIFICATION_TOKEN_REPLACEMENTS.get(token, token)

#         if token and token not in CLASSIFICATION_STOPWORDS:
#             tokens.add(token)

#     return tokens


# def canonicalizar_clasificacion(value: object) -> str:
#     """
#     Agrupa variantes frecuentes de una misma clasificación.

#     Ejemplos equivalentes:
#     - ``EQUIPO EN INDOOR`` / ``Equipo indoor``
#     - ``NÚMERO DE SERIE`` / ``Números de serie``
#     - ``ROTULADO`` / ``Etiquetado`` / ``Rotulado y etiquetado``
#     - ``Condición de herrajes y soportes`` /
#       ``Condición de soporte y herraje``
#     """

#     tokens = _classification_tokens(value)

#     if "outdoor" in tokens:
#         return "equipo outdoor"

#     if "indoor" in tokens:
#         return "equipo indoor"

#     if "conector" in tokens:
#         return "condicion conector"

#     if "marca" in tokens and "modelo" in tokens and "serie" in tokens:
#         return "marca modelo serie"

#     if "serie" in tokens:
#         return "numero serie"

#     if "rotulado" in tokens or "etiquetado" in tokens:
#         return "rotulado etiquetado"

#     if "herraje" in tokens and "soporte" in tokens:
#         return "condicion herraje soporte"

#     if "herraje" in tokens:
#         return "condicion herraje"

#     if "soporte" in tokens:
#         return "condicion soporte"

#     return " ".join(sorted(tokens))


# def _name_similarity(left: object, right: object) -> float:
#     """Calcula una similitud conservadora entre dos nombres de carpeta."""

#     left_normalized = normalizar_texto(limpiar_prefijo_numerico(left))
#     right_normalized = normalizar_texto(limpiar_prefijo_numerico(right))

#     if not left_normalized or not right_normalized:
#         return 0.0

#     if left_normalized == right_normalized:
#         return 1.0

#     left_tokens = _classification_tokens(left)
#     right_tokens = _classification_tokens(right)

#     if left_tokens or right_tokens:
#         union = left_tokens | right_tokens
#         token_score = len(left_tokens & right_tokens) / len(union)
#     else:
#         token_score = 0.0

#     sequence_score = SequenceMatcher(
#         None,
#         left_normalized,
#         right_normalized,
#     ).ratio()

#     return max(token_score, sequence_score)


# def buscar_carpeta_clasificacion(
#     ruta_sitio: str | Path | None,
#     clasificacion_buscada: object,
# ) -> str | None:
#     """
#     Busca una clasificación con varias estrategias, sin depender de que el
#     nombre sea idéntico al del Excel.

#     El orden es:
#     1. Coincidencia flexible exacta.
#     2. Coincidencia canónica de clasificaciones conocidas.
#     3. Coincidencia aproximada con umbral conservador.
#     """

#     exact_match = buscar_carpeta_flexible(
#         ruta_padre=ruta_sitio,
#         nombre_buscado=clasificacion_buscada,
#     )

#     if exact_match:
#         return exact_match

#     if not ruta_sitio or not _fs_is_dir(ruta_sitio):
#         return None

#     expected_canonical = canonicalizar_clasificacion(clasificacion_buscada)
#     candidates = [
#         child
#         for child in _iter_directory(ruta_sitio)
#         if _fs_is_dir(child)
#     ]

#     canonical_matches = [
#         candidate
#         for candidate in candidates
#         if expected_canonical
#         and canonicalizar_clasificacion(candidate.name) == expected_canonical
#     ]

#     if len(canonical_matches) == 1:
#         logger.info(
#             "Clasificación asociada por equivalencia: '%s' -> '%s'.",
#             clasificacion_buscada,
#             canonical_matches[0].name,
#         )
#         return str(canonical_matches[0])

#     scored_candidates = sorted(
#         (
#             (_name_similarity(clasificacion_buscada, candidate.name), candidate)
#             for candidate in candidates
#         ),
#         key=lambda item: item[0],
#         reverse=True,
#     )

#     if not scored_candidates:
#         return None

#     best_score, best_candidate = scored_candidates[0]
#     second_score = scored_candidates[1][0] if len(scored_candidates) > 1 else 0.0

#     if best_score >= 0.62 and (best_score - second_score >= 0.08 or best_score >= 0.90):
#         logger.info(
#             "Clasificación asociada aproximadamente: '%s' -> '%s' (%.2f).",
#             clasificacion_buscada,
#             best_candidate.name,
#             best_score,
#         )
#         return str(best_candidate)

#     return None


# def limpiar_fragmento_nombre_archivo(valor: object) -> str:
#     """Genera un fragmento seguro para nombres de archivos temporales."""

#     texto = str(valor).strip()
#     texto = re.sub(r'[<>:"/\\|?*]', "_", texto)
#     texto = re.sub(r"\s+", "_", texto)
#     texto = texto.strip(" ._")

#     return texto or "SIN_ID"


# def texto_seguro_excel(valor: object) -> str:
#     """Evita que valores vacíos de Excel se conviertan en la cadena ``nan``."""

#     if pd.isna(valor):
#         return ""

#     texto = str(valor).strip()

#     if texto.lower() == "nan":
#         return ""

#     return texto


# def _to_extended_windows_path(path: str | Path) -> str:
#     """
#     Devuelve una ruta absoluta para operaciones de entrada/salida.

#     En Windows agrega el prefijo ``\\?\\`` únicamente al momento de abrir,
#     crear, recorrer o consultar archivos. El resto de la aplicación conserva
#     rutas normales, porque ``pathlib`` y algunas librerías de terceros pueden
#     perder la navegación de carpetas cuando reciben permanentemente una ruta
#     con el prefijo extendido.
#     """

#     absolute_path = os.path.abspath(os.fspath(path))

#     if os.name != "nt":
#         return absolute_path

#     if absolute_path.startswith("\\\\?\\"):
#         return absolute_path

#     if absolute_path.startswith("\\\\"):
#         return "\\\\?\\UNC\\" + absolute_path[2:]

#     return "\\\\?\\" + absolute_path

# def _fs_exists(path: str | Path) -> bool:
#     """Comprueba existencia usando rutas extendidas cuando se ejecuta en Windows."""

#     return os.path.exists(_to_extended_windows_path(path))


# def _fs_is_dir(path: str | Path) -> bool:
#     """Comprueba si una ruta es carpeta sin perder compatibilidad con rutas largas."""

#     return os.path.isdir(_to_extended_windows_path(path))


# def _fs_is_file(path: str | Path) -> bool:
#     """Comprueba si una ruta es archivo sin perder compatibilidad con rutas largas."""

#     return os.path.isfile(_to_extended_windows_path(path))


# def _iter_directory(path: str | Path) -> list[Path]:
#     """
#     Lista el contenido de una carpeta usando ``os.scandir`` sobre la ruta
#     extendida, pero devuelve objetos ``Path`` normales para que la lógica de
#     búsqueda y comparación continúe funcionando igual que antes.
#     """

#     normal_path = Path(path)

#     if not _fs_is_dir(normal_path):
#         return []

#     children: list[Path] = []

#     with os.scandir(_to_extended_windows_path(normal_path)) as entries:
#         for entry in entries:
#             children.append(normal_path / entry.name)

#     return children


# def _walk_directories(root: str | Path) -> list[Path]:
#     """Recorre carpetas de forma compatible con rutas largas de Windows."""

#     root_path = Path(root)
#     directories: list[Path] = []
#     pending: list[Path] = [root_path]

#     while pending:
#         current = pending.pop()

#         for child in _iter_directory(current):
#             if _fs_is_dir(child):
#                 directories.append(child)
#                 pending.append(child)

#     return directories


# def buscar_carpeta_flexible(
#     ruta_padre: str | Path | None,
#     nombre_buscado: object,
# ) -> str | None:
#     """
#     Busca una carpeta ignorando prefijos numéricos, acentos, mayúsculas,
#     guiones y guiones bajos.

#     La navegación se realiza con las funciones de compatibilidad de rutas
#     largas para que el prefijo extendido de Windows no provoque falsos "no existe".
#     """

#     if not ruta_padre:
#         return None

#     ruta_padre_path = Path(ruta_padre)

#     if not _fs_is_dir(ruta_padre_path):
#         return None

#     nombre_buscado_limpio = limpiar_prefijo_numerico(nombre_buscado)
#     nombre_buscado_normalizado = normalizar_texto(nombre_buscado_limpio)

#     if not nombre_buscado_normalizado:
#         return None

#     for carpeta in _iter_directory(ruta_padre_path):
#         if not _fs_is_dir(carpeta):
#             continue

#         carpeta_limpia = limpiar_prefijo_numerico(carpeta.name)

#         if normalizar_texto(carpeta_limpia) == nombre_buscado_normalizado:
#             return str(carpeta)

#     return None

# def _zip_member_parts(member_name: str) -> tuple[str, ...]:
#     """Valida y convierte una ruta interna del ZIP en componentes seguros."""

#     normalized_name = member_name.replace("\\", "/")

#     if "\x00" in normalized_name:
#         raise ValueError("El ZIP contiene un nombre de archivo no permitido.")

#     pure_path = PurePosixPath(normalized_name)

#     if pure_path.is_absolute():
#         raise ValueError("El ZIP contiene rutas absolutas no permitidas.")

#     parts = tuple(part for part in pure_path.parts if part not in ("", "."))

#     if not parts:
#         return ()

#     if any(part == ".." for part in parts):
#         raise ValueError("El ZIP contiene rutas que intentan salir del directorio.")

#     for part in parts:
#         if len(part) > 255:
#             raise ValueError(
#                 "El ZIP contiene un nombre individual mayor a 255 caracteres: "
#                 f"{part[:80]}..."
#             )

#         if any(character in part for character in '<>:"|?*'):
#             raise ValueError(
#                 "El ZIP contiene caracteres no permitidos por Windows en: "
#                 f"{part}"
#             )

#         if any(ord(character) < 32 for character in part):
#             raise ValueError(
#                 "El ZIP contiene caracteres de control no permitidos en: "
#                 f"{part}"
#             )

#         if part.endswith((" ", ".")):
#             raise ValueError(
#                 "El ZIP contiene un nombre que termina en espacio o punto: "
#                 f"{part}"
#             )

#         base_name = part.split(".", 1)[0].upper()
#         if base_name in WINDOWS_RESERVED_NAMES:
#             raise ValueError(
#                 "El ZIP contiene un nombre reservado por Windows: "
#                 f"{part}"
#             )

#     return parts


# def _is_zip_symlink(member: zipfile.ZipInfo) -> bool:
#     """Detecta enlaces simbólicos almacenados en un ZIP creado en Unix."""

#     unix_mode = (member.external_attr >> 16) & 0xFFFF
#     return stat.S_ISLNK(unix_mode)


# def safe_extract_zip(zip_path: Path, extract_dir: Path) -> Path:
#     """
#     Extrae un ZIP de forma segura y compatible con rutas largas de Windows.

#     La ruta que se devuelve es una ruta normal. El prefijo extendido de
#     Windows se aplica únicamente en las operaciones de entrada/salida. Esto
#     evita que la lógica posterior de búsqueda de carpetas pierda las rutas de
#     las evidencias y termine mostrando ``[No encontrada]`` en el documento.
#     """

#     zip_path = Path(zip_path)
#     extract_dir = Path(extract_dir)

#     if not _fs_exists(zip_path):
#         raise FileNotFoundError(f"No existe el archivo ZIP: {zip_path}")

#     if not zipfile.is_zipfile(_to_extended_windows_path(zip_path)):
#         raise ValueError("El archivo de evidencias no es un ZIP válido.")

#     os.makedirs(_to_extended_windows_path(extract_dir), exist_ok=True)

#     try:
#         with zipfile.ZipFile(_to_extended_windows_path(zip_path), "r") as zip_ref:
#             for member in zip_ref.infolist():
#                 if member.flag_bits & 0x1:
#                     raise ValueError(
#                         "El ZIP contiene archivos protegidos con contraseña."
#                     )

#                 if _is_zip_symlink(member):
#                     raise ValueError(
#                         "El ZIP contiene enlaces simbólicos, los cuales no están permitidos."
#                     )

#                 parts = _zip_member_parts(member.filename)

#                 if not parts:
#                     continue

#                 destination = extract_dir.joinpath(*parts)
#                 destination_io = _to_extended_windows_path(destination)

#                 if member.is_dir() or member.filename.endswith(("/", "\\")):
#                     os.makedirs(destination_io, exist_ok=True)
#                     continue

#                 os.makedirs(
#                     _to_extended_windows_path(destination.parent),
#                     exist_ok=True,
#                 )

#                 with zip_ref.open(member, "r") as source_file:
#                     with open(destination_io, "wb") as destination_file:
#                         shutil.copyfileobj(
#                             source_file,
#                             destination_file,
#                             length=1024 * 1024,
#                         )

#     except ValueError:
#         raise

#     except OSError as exc:
#         if getattr(exc, "winerror", None) == 206:
#             raise ValueError(
#                 "No fue posible extraer el ZIP porque contiene una ruta que "
#                 "excede los límites admitidos por el sistema de archivos. "
#                 "La aplicación intentó utilizar rutas extendidas; revise que "
#                 "ningún nombre individual supere 255 caracteres."
#             ) from exc

#         raise ValueError(f"No fue posible extraer el ZIP: {exc}") from exc

#     except zipfile.BadZipFile as exc:
#         raise ValueError(
#             "El archivo de evidencias está dañado o no es un ZIP válido."
#         ) from exc

#     return extract_dir

# def find_evidence_root(
#     extract_dir: Path,
#     expected_link_names: list[str] | tuple[str, ...] | None = None,
# ) -> Path:
#     """
#     Detecta la carpeta que contiene directamente los enlaces del proyecto.

#     Primero compara los nombres de enlace registrados en el Excel contra las
#     carpetas extraídas. Esto es más confiable que buscar cualquier carpeta que
#     contenga la palabra "evidencia", porque dentro del ZIP puede haber carpetas
#     intermedias o clasificaciones con nombres similares.
#     """

#     extract_dir = Path(extract_dir)

#     if not _fs_is_dir(extract_dir):
#         return extract_dir

#     expected_normalized = {
#         normalizar_texto(limpiar_prefijo_numerico(name))
#         for name in (expected_link_names or [])
#         if normalizar_texto(name)
#     }

#     candidate_directories = [extract_dir, *_walk_directories(extract_dir)]

#     if expected_normalized:
#         best_candidate: Path | None = None
#         best_score = 0

#         for candidate in candidate_directories:
#             child_names = {
#                 normalizar_texto(limpiar_prefijo_numerico(child.name))
#                 for child in _iter_directory(candidate)
#                 if _fs_is_dir(child)
#             }

#             score = len(expected_normalized.intersection(child_names))

#             if score > best_score:
#                 best_candidate = candidate
#                 best_score = score

#         if best_candidate is not None and best_score > 0:
#             logger.info(
#                 "Raíz de evidencias detectada por coincidencia con Excel: %s",
#                 best_candidate,
#             )
#             return best_candidate

#     evidence_names = {
#         "evidencia de los enlaces",
#         "evidencias de los enlaces",
#         "evidencia enlaces",
#         "evidencias enlaces",
#     }

#     evidence_candidates = [
#         path
#         for path in candidate_directories
#         if normalizar_texto(path.name) in evidence_names
#     ]

#     if evidence_candidates:
#         evidence_candidates.sort(key=lambda item: len(item.parts))
#         logger.info(
#             "Raíz de evidencias detectada por nombre de carpeta: %s",
#             evidence_candidates[0],
#         )
#         return evidence_candidates[0]

#     current = extract_dir

#     while True:
#         child_directories = [
#             child
#             for child in _iter_directory(current)
#             if _fs_is_dir(child)
#         ]
#         child_files = [
#             child
#             for child in _iter_directory(current)
#             if _fs_is_file(child)
#         ]

#         if len(child_directories) == 1 and not child_files:
#             current = child_directories[0]
#             continue

#         break

#     logger.info("Raíz de evidencias utilizada: %s", current)
#     return current

# def format_excel_date(value: object) -> str:
#     """Convierte fechas de Excel al formato ``dd/mm/aaaa``."""

#     if isinstance(value, (pd.Timestamp, datetime.datetime, datetime.date)):
#         return value.strftime("%d/%m/%Y")

#     if pd.isna(value):
#         return ""

#     return str(value)


# def _strip_supported_image_extension(value: str) -> str:
#     """
#     Elimina únicamente extensiones de imagen conocidas.

#     Es importante no utilizar ``Path.suffix`` directamente para valores como
#     ``1.0`` porque Python interpreta ``.0`` como si fuera una extensión. Para
#     este sistema, ``1`` y ``1.0`` representan el mismo identificador de foto.
#     """

#     text = value.strip()
#     lower_text = text.casefold()

#     for extension in sorted(IMAGE_EXTENSIONS, key=len, reverse=True):
#         if lower_text.endswith(extension):
#             return text[: -len(extension)]

#     return text


# def _normalizar_identificador_imagen(value: object) -> str:
#     """
#     Normaliza identificadores de imagen sin depender del sufijo ``.0``.

#     Los siguientes valores se consideran equivalentes:

#     - ``1``
#     - ``01``
#     - ``001``
#     - ``1.0``
#     - ``01.0``
#     - ``1.jpg``
#     - ``1.0.jpeg``
#     - ``1 - evidencia.jpg``
#     - ``1.0_foto.png``

#     Para nombres no numéricos conserva una comparación textual normalizada.
#     """

#     text = texto_seguro_excel(value)

#     if not text:
#         return ""

#     stem = _strip_supported_image_extension(Path(text).name).strip()

#     numeric_match = re.fullmatch(r"0*(\d+)(?:[.,]0+)?", stem)

#     if numeric_match:
#         return str(int(numeric_match.group(1)))


#     numeric_prefix_match = re.match(
#         r"^\s*0*(\d+)(?:[.,]0+)?(?=$|[\s_-])",
#         stem,
#     )

#     if numeric_prefix_match:
#         return str(int(numeric_prefix_match.group(1)))

#     return normalizar_texto(stem)

# def _iter_files_recursive(
#     root: str | Path | None,
#     *,
#     max_depth: int = 6,
# ) -> list[Path]:
#     """Lista archivos de manera recursiva y compatible con rutas largas."""

#     if not root or not _fs_is_dir(root):
#         return []

#     root_path = Path(root)
#     files: list[Path] = []
#     pending: list[tuple[Path, int]] = [(root_path, 0)]

#     while pending:
#         current, depth = pending.pop()

#         for child in _iter_directory(current):
#             if _fs_is_file(child):
#                 files.append(child)
#             elif _fs_is_dir(child) and depth < max_depth:
#                 pending.append((child, depth + 1))

#     return files


# def _image_name_matches(file_path: Path, requested_name: object) -> bool:
#     """
#     Compara el nombre solicitado por el Excel con una imagen real.

#     La comparación no depende de que el archivo se llame ``1.0.jpg``. También
#     reconoce ``1.jpg``, ``01.jpeg``, ``001.png`` y nombres descriptivos que
#     comienzan con el mismo identificador numérico.
#     """

#     if file_path.suffix.lower() not in IMAGE_EXTENSIONS:
#         return False

#     requested_text = texto_seguro_excel(requested_name)

#     if not requested_text:
#         return False

#     requested_base = _strip_supported_image_extension(
#         Path(requested_text).name
#     ).strip()
#     file_base = _strip_supported_image_extension(file_path.name).strip()

#     if requested_base.casefold() == file_base.casefold():
#         return True

#     requested_id = _normalizar_identificador_imagen(requested_text)
#     file_id = _normalizar_identificador_imagen(file_path.name)

#     if requested_id and requested_id == file_id:
#         return True

    
#     if requested_id.isdigit():
#         numeric_prefix = re.match(
#             r"^\s*0*(\d+)(?:[.,]0+)?(?=$|[\s_-])",
#             file_base,
#             flags=re.IGNORECASE,
#         )

#         if numeric_prefix:
#             return str(int(numeric_prefix.group(1))) == requested_id

#     return False



# def buscar_imagen_flexible(
#     ruta_carpeta: str | Path | None,
#     nombre_evidencia: object,
# ) -> Path | None:
#     """
#     Busca una imagen dentro de la clasificación y sus subcarpetas.

#     Esta búsqueda acepta diferencias de extensión, mayúsculas, valores
#     numéricos provenientes de Excel y carpetas intermedias como ``Fotos``.
#     """

#     if not ruta_carpeta or not _fs_is_dir(ruta_carpeta):
#         return None

#     matches = [
#         file_path
#         for file_path in _iter_files_recursive(ruta_carpeta)
#         if _image_name_matches(file_path, nombre_evidencia)
#     ]

#     if len(matches) == 1:
#         return matches[0]

#     if len(matches) > 1:
#         # Preferir el archivo menos profundo y con nombre más corto.
#         matches.sort(
#             key=lambda item: (
#                 len(item.relative_to(Path(ruta_carpeta)).parts),
#                 len(item.name),
#                 item.name.casefold(),
#             )
#         )
#         return matches[0]

#     return None


# def buscar_imagen_contextual(
#     *,
#     ruta_clasificacion: str | Path | None,
#     ruta_sitio: str | Path | None,
#     ruta_enlace: str | Path | None,
#     clasificacion: object,
#     nombre_evidencia: object,
# ) -> Path | None:
#     """
#     Localiza una evidencia aunque el ZIP tenga una carpeta extra o una
#     variante razonable en el nombre de la clasificación.

#     Nunca selecciona una coincidencia global ambigua. La prioridad es:
#     1. Carpeta de clasificación resuelta.
#     2. Candidatos bajo el sitio, puntuados por su carpeta padre.
#     3. Candidato único bajo el enlace.
#     """

#     direct_match = buscar_imagen_flexible(
#         ruta_carpeta=ruta_clasificacion,
#         nombre_evidencia=nombre_evidencia,
#     )

#     if direct_match:
#         return direct_match

#     site_candidates = [
#         file_path
#         for file_path in _iter_files_recursive(ruta_sitio)
#         if _image_name_matches(file_path, nombre_evidencia)
#     ]

#     if len(site_candidates) == 1:
#         logger.info(
#             "Imagen '%s' encontrada mediante búsqueda recursiva del sitio: %s",
#             nombre_evidencia,
#             site_candidates[0],
#         )
#         return site_candidates[0]

#     if len(site_candidates) > 1:
#         expected_canonical = canonicalizar_clasificacion(clasificacion)
#         scored: list[tuple[float, Path]] = []

#         for candidate in site_candidates:
#             # Se revisan hasta tres carpetas ascendentes porque algunas
#             # estructuras agregan carpetas intermedias como ``Fotos``.
#             ancestors = list(candidate.parents)[:3]
#             scores = []

#             for ancestor in ancestors:
#                 canonical_score = (
#                     1.0
#                     if expected_canonical
#                     and canonicalizar_clasificacion(ancestor.name) == expected_canonical
#                     else 0.0
#                 )
#                 scores.append(
#                     max(
#                         canonical_score,
#                         _name_similarity(clasificacion, ancestor.name),
#                     )
#                 )

#             scored.append((max(scores, default=0.0), candidate))

#         scored.sort(key=lambda item: item[0], reverse=True)
#         best_score, best_candidate = scored[0]
#         second_score = scored[1][0] if len(scored) > 1 else 0.0

#         if best_score >= 0.62 and (best_score - second_score >= 0.08 or best_score >= 0.95):
#             logger.info(
#                 "Imagen '%s' asociada por contexto de clasificación '%s': %s",
#                 nombre_evidencia,
#                 clasificacion,
#                 best_candidate,
#             )
#             return best_candidate

#     link_candidates = [
#         file_path
#         for file_path in _iter_files_recursive(ruta_enlace)
#         if _image_name_matches(file_path, nombre_evidencia)
#     ]

#     if len(link_candidates) == 1:
#         logger.info(
#             "Imagen '%s' encontrada como coincidencia única del enlace: %s",
#             nombre_evidencia,
#             link_candidates[0],
#         )
#         return link_candidates[0]

#     return None

# def normalize_image_for_docx(
#     source_path: Path,
#     normalized_images_dir: Path,
# ) -> Path:
#     """
#     Corrige orientación, elimina metadatos problemáticos y genera un PNG limpio.
#     """

#     source_path = Path(source_path)
#     normalized_images_dir = Path(normalized_images_dir)

#     if not _fs_exists(source_path):
#         raise InvalidEvidenceImageError(
#             f"No se encontró la imagen: {source_path.name}"
#         )

#     if not _fs_is_file(source_path):
#         raise InvalidEvidenceImageError(
#             f"La ruta no corresponde a una imagen: {source_path}"
#         )

#     normalized_images_dir.mkdir(parents=True, exist_ok=True)
#     normalized_path = normalized_images_dir / (
#         f"{limpiar_fragmento_nombre_archivo(source_path.stem)}_"
#         f"{uuid4().hex}.png"
#     )

#     try:
#         with Image.open(_to_extended_windows_path(source_path)) as source_image:
#             source_image.load()
#             corrected_image = ImageOps.exif_transpose(source_image)

#             if corrected_image.mode in ("RGBA", "LA"):
#                 normalized_image = corrected_image.convert("RGBA")
#             elif corrected_image.mode == "P":
#                 if "transparency" in corrected_image.info:
#                     normalized_image = corrected_image.convert("RGBA")
#                 else:
#                     normalized_image = corrected_image.convert("RGB")
#             else:
#                 normalized_image = corrected_image.convert("RGB")

#             normalized_image.save(
#                 normalized_path,
#                 format="PNG",
#                 optimize=True,
#             )

#     except UnidentifiedImageError as exc:
#         raise InvalidEvidenceImageError(
#             "El archivo no contiene una imagen válida: "
#             f"{source_path.name}"
#         ) from exc

#     except (OSError, ValueError, SyntaxError) as exc:
#         raise InvalidEvidenceImageError(
#             "La imagen está dañada, incompleta o utiliza un formato "
#             f"incompatible: {source_path.name}"
#         ) from exc

#     if not normalized_path.exists() or normalized_path.stat().st_size == 0:
#         raise InvalidEvidenceImageError(
#             f"No fue posible preparar la imagen: {source_path.name}"
#         )

#     return normalized_path


# def build_inline_image(
#     template: DocxTemplate,
#     image_path: str | Path,
#     normalized_images_dir: Path,
# ) -> InlineImage:
#     """Normaliza una imagen y crea el ``InlineImage`` para docxtpl."""

#     normalized_path = normalize_image_for_docx(
#         source_path=Path(image_path),
#         normalized_images_dir=normalized_images_dir,
#     )

#     try:
#         with Image.open(normalized_path) as image:
#             width_pixels, height_pixels = image.size
#     except Exception as exc:
#         raise InvalidEvidenceImageError(
#             "No fue posible leer la imagen normalizada: "
#             f"{Path(image_path).name}"
#         ) from exc

#     if width_pixels >= height_pixels:
#         return InlineImage(template, str(normalized_path), width=Mm(70))

#     return InlineImage(template, str(normalized_path), height=Mm(52.5))


# def generate_protocol_document(
#     excel_path: Path,
#     evidence_zip_path: Path,
#     template_docx_path: Path,
#     work_dir: Path,
# ) -> ProtocolDocumentResult:
#     """Genera el protocolo Word final a partir del Excel y ZIP de evidencias."""

#     excel_path = Path(excel_path)
#     evidence_zip_path = Path(evidence_zip_path)
#     template_docx_path = Path(template_docx_path)
#     work_dir = Path(work_dir)

#     if not excel_path.exists():
#         raise FileNotFoundError(f"No existe el Excel: {excel_path}")

#     if not evidence_zip_path.exists():
#         raise FileNotFoundError(
#             f"No existe el ZIP de evidencias: {evidence_zip_path}"
#         )

#     if not template_docx_path.exists():
#         raise FileNotFoundError(
#             f"No existe la plantilla Word: {template_docx_path}"
#         )

#     job_id = uuid4().hex
#     job_dir = work_dir / job_id

#     # Nombres deliberadamente cortos. Además del soporte de rutas extendidas,
#     # esto reduce la longitud total y mejora compatibilidad con herramientas
#     # de terceros que todavía usan MAX_PATH.
#     extract_dir = job_dir / "e"
#     normalized_images_dir = job_dir / "i"
#     temp_dir = job_dir / "t"
#     output_dir = job_dir / "o"
#     output_docx_path = output_dir / "Protocolo_C5_Generado.docx"

#     if job_dir.exists():
#         shutil.rmtree(job_dir, ignore_errors=True)

#     extract_dir.mkdir(parents=True, exist_ok=True)
#     normalized_images_dir.mkdir(parents=True, exist_ok=True)
#     temp_dir.mkdir(parents=True, exist_ok=True)
#     output_dir.mkdir(parents=True, exist_ok=True)

#     effective_extract_dir = safe_extract_zip(
#         zip_path=evidence_zip_path,
#         extract_dir=extract_dir,
#     )

#     try:
#         hoja1 = pd.read_excel(excel_path, sheet_name=0)
#         hoja2 = pd.read_excel(excel_path, sheet_name=1)
#     except Exception as exc:
#         raise ValueError(
#             "No fue posible leer la plantilla Excel. Verifique que el archivo "
#             "corresponda a la plantilla oficial y que no se encuentre dañado."
#         ) from exc

#     hoja1.columns = hoja1.columns.astype(str).str.strip()
#     hoja2.columns = hoja2.columns.astype(str).str.strip()

#     required_hoja1 = ["ID", "Sitio"]
#     required_hoja2 = ["ID", "Sitio", "clasificacion", "Foto"]

#     missing_hoja1 = [
#         column for column in required_hoja1 if column not in hoja1.columns
#     ]
#     missing_hoja2 = [
#         column for column in required_hoja2 if column not in hoja2.columns
#     ]

#     if missing_hoja1:
#         raise ValueError(
#             "La primera hoja del Excel no contiene las columnas requeridas: "
#             + ", ".join(missing_hoja1)
#         )

#     if missing_hoja2:
#         raise ValueError(
#             "La segunda hoja del Excel no contiene las columnas requeridas: "
#             + ", ".join(missing_hoja2)
#         )

#     hoja1["ID"] = hoja1["ID"].astype(str).str.strip()
#     hoja2["ID"] = hoja2["ID"].astype(str).str.strip()

#     expected_link_names = [
#         texto_seguro_excel(value)
#         for value in hoja1.get("Enlace", pd.Series(dtype=object)).tolist()
#         if texto_seguro_excel(value)
#     ]

#     ruta_evidencias = find_evidence_root(
#         extract_dir=effective_extract_dir,
#         expected_link_names=expected_link_names,
#     )

#     archivos_generados: list[Path] = []
#     total_sitios = len(hoja1)
#     total_imagenes_omitidas = 0
#     total_imagenes_insertadas = 0
#     total_imagenes_no_encontradas = 0
#     total_clasificaciones_no_encontradas = 0

#     imagenes_no_encontradas: list[dict[str, str]] = []
#     imagenes_omitidas: list[dict[str, str]] = []
#     clasificaciones_no_encontradas: list[dict[str, str]] = []
#     clasificaciones_no_encontradas_keys: set[tuple[str, str]] = set()

#     total_imagenes_en_zip = sum(
#         1
#         for file_path in _iter_files_recursive(ruta_evidencias, max_depth=12)
#         if file_path.suffix.lower() in IMAGE_EXTENSIONS
#     )

#     logger.info(
#         "Diagnóstico del ZIP: raíz='%s', imágenes detectadas=%s.",
#         ruta_evidencias,
#         total_imagenes_en_zip,
#     )

#     if total_imagenes_en_zip == 0:
#         raise ValueError(
#             "El ZIP de evidencias no contiene imágenes compatibles. "
#             "Primero extraiga el ZIP de estructura, coloque las evidencias "
#             "dentro de las carpetas correspondientes y vuelva a comprimirlo."
#         )

#     if total_sitios == 0:
#         raise ValueError("La primera hoja del Excel no contiene registros.")

#     for index, fila_sitio in hoja1.iterrows():
#         id_sitio = texto_seguro_excel(fila_sitio["ID"])
#         nombre_sitio_real = texto_seguro_excel(fila_sitio["Sitio"])
#         nombre_enlace = texto_seguro_excel(fila_sitio.get("Enlace", ""))
#         tipo_sitio_original = texto_seguro_excel(
#             fila_sitio.get("TipoDeSitio", "")
#         ).upper()

#         if "LOCAL" in tipo_sitio_original:
#             tipo_sitio_formateado = "Sitio Local"
#         else:
#             tipo_sitio_formateado = "Sitio Remoto"

#         template = DocxTemplate(str(template_docx_path))

#         registros_fotos = hoja2[
#             (hoja2["ID"] == id_sitio)
#             & (
#                 hoja2["Sitio"].astype(str).str.strip()
#                 == nombre_sitio_real
#             )
#         ]

#         clasificaciones_unicas = (
#             registros_fotos["clasificacion"]
#             .fillna("SIN CLASIFICACION")
#             .astype(str)
#             .str.strip()
#             .unique()
#         )

#         lista_clasificaciones = []

#         ruta_enlace_folder = buscar_carpeta_flexible(
#             ruta_padre=ruta_evidencias,
#             nombre_buscado=nombre_enlace,
#         )
#         ruta_sitio_folder = buscar_carpeta_flexible(
#             ruta_padre=ruta_enlace_folder,
#             nombre_buscado=nombre_sitio_real,
#         )

#         if not ruta_enlace_folder:
#             logger.warning(
#                 "No se encontró la carpeta del enlace '%s' dentro de '%s'.",
#                 nombre_enlace,
#                 ruta_evidencias,
#             )

#         elif not ruta_sitio_folder:
#             logger.warning(
#                 "No se encontró la carpeta del sitio '%s' dentro del enlace '%s'.",
#                 nombre_sitio_real,
#                 ruta_enlace_folder,
#             )

#         for clasificacion in clasificaciones_unicas:
#             fotos_clase = registros_fotos[
#                 (
#                     registros_fotos["clasificacion"]
#                     .fillna("SIN CLASIFICACION")
#                     .astype(str)
#                     .str.strip()
#                 )
#                 == clasificacion
#             ]

#             ruta_clasificacion = buscar_carpeta_clasificacion(
#                 ruta_sitio=ruta_sitio_folder,
#                 clasificacion_buscada=clasificacion,
#             )

#             if not ruta_clasificacion:
#                 logger.warning(
#                     "No se encontró la clasificación '%s' para el sitio '%s'.",
#                     clasificacion,
#                     nombre_sitio_real,
#                 )

#                 classification_key = (
#                     nombre_sitio_real.casefold(),
#                     str(clasificacion).casefold(),
#                 )

#                 if classification_key not in clasificaciones_no_encontradas_keys:
#                     clasificaciones_no_encontradas_keys.add(classification_key)
#                     total_clasificaciones_no_encontradas += 1
#                     clasificaciones_no_encontradas.append(
#                         {
#                             "sitio": nombre_sitio_real,
#                             "clasificacion": str(clasificacion),
#                             "enlace": nombre_enlace,
#                         }
#                     )

#             lista_individual = []

#             for _, fila_foto in fotos_clase.iterrows():
#                 nombre_evidencia = texto_seguro_excel(
#                     fila_foto.get("Foto", "")
#                 )
#                 nombre_documento = texto_seguro_excel(
#                     fila_foto.get("NombreDeImagenEnDocumento", "")
#                 )

#                 if not nombre_documento:
#                     nombre_documento = " "

#                 datos_bloque = {
#                     "nombre_foto": nombre_documento,
#                     "tipo_equipo": texto_seguro_excel(
#                         fila_foto.get("TipoDeEquipo", "")
#                     ),
#                     "sitio": nombre_sitio_real,
#                     "imagen": "",
#                 }

#                 ruta_imagen = buscar_imagen_contextual(
#                     ruta_clasificacion=ruta_clasificacion,
#                     ruta_sitio=ruta_sitio_folder,
#                     ruta_enlace=ruta_enlace_folder,
#                     clasificacion=clasificacion,
#                     nombre_evidencia=nombre_evidencia,
#                 )

#                 if ruta_imagen:
#                     try:
#                         datos_bloque["imagen"] = build_inline_image(
#                             template=template,
#                             image_path=ruta_imagen,
#                             normalized_images_dir=normalized_images_dir,
#                         )
#                         total_imagenes_insertadas += 1
#                     except InvalidEvidenceImageError as exc:
#                         total_imagenes_omitidas += 1
#                         logger.warning(
#                             "La imagen '%s' fue omitida. Motivo: %s",
#                             ruta_imagen,
#                             exc,
#                         )
#                         imagenes_omitidas.append(
#                             {
#                                 "sitio": nombre_sitio_real,
#                                 "clasificacion": str(clasificacion),
#                                 "imagen_solicitada": nombre_evidencia,
#                                 "archivo_detectado": ruta_imagen.name,
#                                 "motivo": str(exc),
#                             }
#                         )
#                         datos_bloque["imagen"] = (
#                             "[Imagen no disponible: "
#                             f"{nombre_evidencia}]"
#                         )
#                 else:
#                     total_imagenes_no_encontradas += 1
#                     logger.warning(
#                         "No se encontró la imagen '%s' para sitio='%s', "
#                         "clasificación='%s'. Carpeta de clasificación=%s.",
#                         nombre_evidencia,
#                         nombre_sitio_real,
#                         clasificacion,
#                         ruta_clasificacion,
#                     )

#                     imagenes_no_encontradas.append(
#                         {
#                             "sitio": nombre_sitio_real,
#                             "clasificacion": str(clasificacion),
#                             "imagen_solicitada": nombre_evidencia,
#                             "enlace": nombre_enlace,
#                             "motivo": (
#                                 "No se encontró la carpeta de clasificación."
#                                 if not ruta_clasificacion
#                                 else "No existe un archivo compatible dentro de la clasificación."
#                             ),
#                         }
#                     )

#                     datos_bloque["imagen"] = (
#                         f"[No encontrada: {nombre_evidencia}]"
#                     )

#                 lista_individual.append(datos_bloque)

#             lista_pares = []

#             for position in range(0, len(lista_individual), 2):
#                 pair = {
#                     "izq": lista_individual[position],
#                     "der": (
#                         lista_individual[position + 1]
#                         if position + 1 < len(lista_individual)
#                         else None
#                     ),
#                 }
#                 lista_pares.append(pair)

#             lista_clasificaciones.append(
#                 {
#                     "clasificacion": clasificacion,
#                     "pares": lista_pares,
#                 }
#             )

#         fecha_arribo = format_excel_date(
#             fila_sitio.get("FechaDeArribo", "")
#         )
#         fecha_terminacion = format_excel_date(
#             fila_sitio.get("FechaDeTerminación", "")
#         )

#         datos_render = {
#             "ID": id_sitio,
#             "Enlace": nombre_enlace,
#             "Sitio": nombre_sitio_real,
#             "NoTicket": texto_seguro_excel(fila_sitio.get("NoTicket", "")),
#             "TipoDeSitio": tipo_sitio_formateado,
#             "TipoDeEquipo": texto_seguro_excel(
#                 fila_sitio.get("TipoDeEquipo", "")
#             ),
#             "Dato1": fecha_arribo,
#             "Dato2": texto_seguro_excel(
#                 fila_sitio.get("HoraDeArribo", "")
#             ),
#             "Dato3": fecha_terminacion,
#             "Dato4": texto_seguro_excel(
#                 fila_sitio.get("HoraDeTerminación", "")
#             ),
#             "clasificaciones": lista_clasificaciones,
#         }

#         template.render(datos_render)

#         safe_id = limpiar_fragmento_nombre_archivo(id_sitio)
#         temporal_path = temp_dir / f"temp_sitio_{safe_id}_{index}.docx"
#         template.save(str(temporal_path))
#         archivos_generados.append(temporal_path)

#     logger.info(
#         "Resumen de imágenes: detectadas_en_zip=%s, insertadas=%s, "
#         "no_encontradas=%s, omitidas_por_daño=%s.",
#         total_imagenes_en_zip,
#         total_imagenes_insertadas,
#         total_imagenes_no_encontradas,
#         total_imagenes_omitidas,
#     )

#     diagnostics_path = output_dir / "diagnostico_generacion.json"
#     diagnostics_payload = {
#         "schema_version": "1.0",
#         "job_id": job_id,
#         "estado": "completado_con_observaciones" if (
#             total_imagenes_no_encontradas
#             or total_imagenes_omitidas
#             or total_clasificaciones_no_encontradas
#         ) else "completado",
#         "resumen": {
#             "imagenes_detectadas_en_zip": total_imagenes_en_zip,
#             "imagenes_insertadas": total_imagenes_insertadas,
#             "imagenes_no_encontradas": total_imagenes_no_encontradas,
#             "imagenes_omitidas_por_dano": total_imagenes_omitidas,
#             "clasificaciones_no_encontradas": total_clasificaciones_no_encontradas,
#         },
#         "imagenes_no_encontradas": imagenes_no_encontradas,
#         "imagenes_omitidas": imagenes_omitidas,
#         "clasificaciones_no_encontradas": clasificaciones_no_encontradas,
#     }

#     diagnostics_path.write_text(
#         json.dumps(
#             diagnostics_payload,
#             ensure_ascii=False,
#             indent=2,
#         ),
#         encoding="utf-8",
#     )

#     if total_imagenes_insertadas == 0:
#         raise ValueError(
#             "El ZIP sí contiene imágenes, pero ninguna pudo asociarse con la "
#             "plantilla Excel. Verifique que esté utilizando el mismo Excel "
#             "con el que se generó la estructura y que las evidencias conserven "
#             "los nombres indicados en la columna 'Foto'."
#         )

#     if not archivos_generados:
#         raise ValueError("No se generaron documentos.")

#     documento_base = Document(str(archivos_generados[0]))
#     composer = Composer(documento_base)

#     for archivo in archivos_generados[1:]:
#         documento_base.add_page_break()
#         documento_temporal = Document(str(archivo))
#         composer.append(documento_temporal)

#     composer.save(str(output_docx_path))

#     if not output_docx_path.exists() or output_docx_path.stat().st_size == 0:
#         raise RuntimeError("El documento Word no se generó correctamente.")

#     return ProtocolDocumentResult(
#         job_id=job_id,
#         output_docx_path=output_docx_path,
#         total_sitios=total_sitios,
#         total_documentos_generados=len(archivos_generados),
#         total_imagenes_omitidas=total_imagenes_omitidas,
#         total_imagenes_no_encontradas=total_imagenes_no_encontradas,
#         total_clasificaciones_no_encontradas=(
#             total_clasificaciones_no_encontradas
#         ),
#         diagnostics_path=diagnostics_path,
#     )









